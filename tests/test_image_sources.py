"""Images pasted into a Conversation become Library Resources, Course Sources, and Slides."""

import json
from collections.abc import AsyncIterator
from io import BytesIO
from pathlib import Path

import httpx2
import pytest
from PIL import Image
from pptx import Presentation as PPTXPresentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pydantic_ai.messages import (
    BinaryContent,
    ModelMessage,
    ModelRequest,
    ToolReturnPart,
    UserPromptPart,
)
from pydantic_ai.models.function import AgentInfo, DeltaToolCall, FunctionModel

from course_harness import resources as res
from course_harness.app import create_app
from course_harness.chat_history import read_chat_history, save_chat_history
from course_harness.course_agent import (
    CourseAgentDeps,
    CourseAgentState,
    ReplacePresentationCommand,
    SlideCommand,
    _build_course_agent,
    apply_presentation_command,
)
from course_harness.course_plan import (
    CoursePlanInput,
    LectureInput,
    create_course_plan,
    create_course_plan_file,
    initialize_workspace_history,
)
from course_harness.export import export_presentation
from course_harness.library import register_and_snapshot
from course_harness.presentation import ImageSlide, Presentation
from course_harness.sources import admit_source, source_image_resolver
from course_harness.workspace_history import record_app_authored_state


def _png(width: int = 40, height: int = 20) -> bytes:
    buffer = BytesIO()
    Image.new("RGB", (width, height), (200, 30, 30)).save(buffer, format="PNG")
    return buffer.getvalue()


def _register_image(data_dir: Path, cache_dir: Path, content: bytes | None = None):
    request = res.ResourceRegistrationRequest(
        kind="upload", location="diagram.png", media_type="image/png"
    )
    return register_and_snapshot(data_dir, cache_dir, request, content or _png())


def _course(workspace: Path):
    workspace.mkdir()
    plan = create_course_plan(
        CoursePlanInput(title="Images", audience="Test", lectures=[LectureInput(title="One")])
    )
    initialize_workspace_history(workspace)
    create_course_plan_file(workspace, plan)
    record_app_authored_state(workspace)
    return plan


def test_image_resource_is_processed_into_a_textual_reference(tmp_path: Path) -> None:
    _, snapshot, state = _register_image(tmp_path / "data", tmp_path / "cache", _png(40, 20))

    assert state.status == "ready"
    reference = tmp_path / "cache" / "derived" / snapshot.content_hash / "extracted.md"
    assert "image/png" in reference.read_text(encoding="utf-8")
    assert "40 × 20 px" in reference.read_text(encoding="utf-8")


def test_unreadable_image_resource_fails_processing(tmp_path: Path) -> None:
    _, _, state = _register_image(tmp_path / "data", tmp_path / "cache", b"not an image")

    assert state.status == "failed"
    assert state.error == "The image could not be read."


def test_export_fits_an_image_source_into_the_picture_placeholder(tmp_path: Path) -> None:
    workspace = tmp_path / "course"
    plan = _course(workspace)
    data_dir, cache_dir = tmp_path / "data", tmp_path / "cache"
    resource, _, _ = _register_image(data_dir, cache_dir, _png(400, 100))
    source = admit_source(workspace, data_dir, cache_dir, resource.id)
    presentation = Presentation(
        id="presentation-abc123def456",
        lecture_id=plan.lectures[0].id,
        slides=[
            ImageSlide(
                id="slide-abc123def456",
                title="Pipeline",
                image_source_id=source.id,
                caption="The data pipeline",
            )
        ],
    )

    exported = PPTXPresentation(
        BytesIO(
            export_presentation(
                presentation, image_resolver=source_image_resolver(workspace, data_dir)
            )
        )
    )

    slide = exported.slides[0]
    pictures = [
        shape
        for shape in slide.shapes
        if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.PLACEHOLDER)
        and hasattr(shape, "image")
    ]
    assert len(pictures) == 1
    picture = pictures[0]
    assert picture.image.blob == (data_dir / "snapshots" / source.source_version_id).read_bytes()
    assert (picture.crop_left, picture.crop_right) == (0, 0)
    assert picture.width / picture.height == pytest.approx(4, rel=0.01)
    texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
    assert "The data pipeline" in texts


def test_presentation_command_rejects_an_image_that_is_not_an_image_source(
    tmp_path: Path,
) -> None:
    workspace = tmp_path / "course"
    plan = _course(workspace)
    data_dir, cache_dir = tmp_path / "data", tmp_path / "cache"
    notes = register_and_snapshot(
        data_dir,
        cache_dir,
        res.ResourceRegistrationRequest(
            kind="upload", location="notes.md", media_type="text/markdown"
        ),
        b"# Notes",
    )[0]
    text_source = admit_source(workspace, data_dir, cache_dir, notes.id)
    command = ReplacePresentationCommand(
        lecture_id=plan.lectures[0].id,
        slides=[SlideCommand(layout="image", image_source_id=text_source.id)],
    )

    with pytest.raises(ValueError, match="not an image Source"):
        apply_presentation_command(
            workspace,
            command,
            plan,
            image_resolver=source_image_resolver(workspace, data_dir),
        )


@pytest.mark.anyio
async def test_course_author_places_an_uploaded_image_on_a_slide(tmp_path: Path) -> None:
    workspace = tmp_path / "course"
    plan = _course(workspace)
    lecture_id = plan.lectures[0].id
    app = create_app(
        workspace,
        library_data_path=tmp_path / "data",
        library_cache_path=tmp_path / "cache",
    )
    transport = httpx2.ASGITransport(app=app)
    async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
        uploaded = await client.post(
            "/api/resources/upload",
            files={"file": ("pasted.png", _png(), "image/png")},
        )
        assert uploaded.status_code == 201
        assert uploaded.json()["status"] == "ready"
        resource_id = uploaded.json()["resource_id"]
        source = (await client.post("/api/sources", json={"resource_id": resource_id})).json()
        created = await client.post(
            f"/api/presentations/{lecture_id}",
            json={"slides": [{"layout": "image", "title": "Figure"}]},
        )
        slide_id = created.json()["slides"][0]["id"]

        placed = await client.patch(
            f"/api/presentations/{lecture_id}/slides/{slide_id}",
            json={"image_source_id": source["id"]},
        )
        image = await client.get(f"/api/sources/{source['id']}/image")
        cleared = await client.patch(
            f"/api/presentations/{lecture_id}/slides/{slide_id}",
            json={"image_source_id": ""},
        )

    assert placed.status_code == 200
    assert placed.json()["slides"][0]["image_source_id"] == source["id"]
    assert image.status_code == 200
    assert image.headers["content-type"] == "image/png"
    assert image.content == _png()
    assert cleared.json()["slides"][0]["image_source_id"] is None


@pytest.mark.anyio
async def test_slide_rejects_a_source_that_is_not_an_image(tmp_path: Path) -> None:
    workspace = tmp_path / "course"
    plan = _course(workspace)
    lecture_id = plan.lectures[0].id
    app = create_app(
        workspace,
        library_data_path=tmp_path / "data",
        library_cache_path=tmp_path / "cache",
    )
    transport = httpx2.ASGITransport(app=app)
    async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
        uploaded = await client.post(
            "/api/resources/upload",
            files={"file": ("notes.md", b"# Notes", "text/markdown")},
        )
        source = (
            await client.post("/api/sources", json={"resource_id": uploaded.json()["resource_id"]})
        ).json()
        created = await client.post(
            f"/api/presentations/{lecture_id}",
            json={"slides": [{"layout": "image", "title": "Figure"}]},
        )
        slide_id = created.json()["slides"][0]["id"]

        rejected = await client.patch(
            f"/api/presentations/{lecture_id}/slides/{slide_id}",
            json={"image_source_id": source["id"]},
        )
        image = await client.get(f"/api/sources/{source['id']}/image")

    assert rejected.status_code == 422
    assert image.status_code == 404


def _view_image_model(image_id: str, seen: list[list[ModelMessage]]) -> FunctionModel:
    async def stream_function(
        messages: list[ModelMessage], _info: AgentInfo
    ) -> AsyncIterator[str | dict[int, DeltaToolCall]]:
        seen.append(list(messages))
        viewed = any(
            isinstance(part, ToolReturnPart) and part.tool_name == "view_image"
            for message in messages
            if isinstance(message, ModelRequest)
            for part in message.parts
        )
        if not viewed:
            yield {
                0: DeltaToolCall(
                    name="view_image",
                    tool_call_id="view-1",
                    json_args=json.dumps({"image_id": image_id}),
                )
            }
            return
        yield "It is a red rectangle."

    return FunctionModel(stream_function=stream_function)


def _images_in(messages: list[ModelMessage]) -> list[BinaryContent]:
    return [
        item
        for message in messages
        if isinstance(message, ModelRequest)
        for part in message.parts
        if isinstance(part, UserPromptPart) and not isinstance(part.content, str)
        for item in part.content
        if isinstance(item, BinaryContent)
    ]


@pytest.mark.anyio
@pytest.mark.parametrize("vision", [True, False])
async def test_agent_views_an_attached_image_only_with_vision(tmp_path: Path, vision: bool) -> None:
    data_dir, cache_dir = tmp_path / "data", tmp_path / "cache"
    resource, _, _ = _register_image(data_dir, cache_dir)
    seen: list[list[ModelMessage]] = []

    agent = _build_course_agent(requires_approval=False)
    async with agent.run_stream(
        f'[Attachment: {resource.id} image/png "diagram.png"]\n\nWhat is in this image?',
        deps=CourseAgentDeps(
            course_state=CourseAgentState(),
            workspace=tmp_path,
            data_dir=data_dir,
            cache_dir=cache_dir,
            vision=vision,
        ),
        model=_view_image_model(resource.id, seen),
    ) as streamed:
        await streamed.get_output()

    final_request = seen[-1]
    tool_return = next(
        part
        for message in final_request
        if isinstance(message, ModelRequest)
        for part in message.parts
        if isinstance(part, ToolReturnPart) and part.tool_name == "view_image"
    )
    assert "40 × 20 px" in str(tool_return.content)
    images = _images_in(final_request)
    if vision:
        assert [image.data for image in images] == [_png()]
    else:
        assert images == []
        assert "cannot see images" in str(tool_return.content)


def test_stored_conversations_keep_image_references_instead_of_bytes(tmp_path: Path) -> None:
    store, workspace = tmp_path / "chat", tmp_path / "course"
    workspace.mkdir()
    history: list[ModelMessage] = [
        ModelRequest(
            parts=[
                UserPromptPart(
                    content=[
                        "Image resource-0123456789ab (diagram.png):",
                        BinaryContent(
                            data=_png(), media_type="image/png", identifier="resource-0123456789ab"
                        ),
                    ]
                )
            ]
        )
    ]

    save_chat_history(store, workspace, history)

    stored = read_chat_history(store, workspace)
    assert _images_in(stored) == []
    part = stored[0].parts[0]
    assert isinstance(part, UserPromptPart)
    assert "resource-0123456789ab was viewed earlier" in str(part.content)
