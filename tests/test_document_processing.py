import io
import json
from pathlib import Path
from unittest.mock import patch

import httpx2
import pytest
from PIL import Image, ImageDraw
from pptx import Presentation as PPTXPresentation

from course_harness import hosted_processors
from course_harness import sources as sources_module
from course_harness.app import create_app
from course_harness.docling_parser import PDF_MEDIA_TYPE, PPTX_MEDIA_TYPE, PageContent
from course_harness.document_processing_settings import DocumentProcessingStore
from course_harness.document_processors import (
    Conversion,
    DocumentProcessor,
    DocumentProcessors,
    check_quality,
    markdown_pages,
)
from course_harness.hosted_processors import (
    LlamaParseProcessor,
    MistralOcrProcessor,
    installed_processors,
)
from course_harness.search import search_raw

FIXTURES = Path(__file__).with_name("fixtures") / "resources"
TEACHING_PDF = (FIXTURES / "teaching.pdf").read_bytes()


async def _public_host_resolver(_hostname: str, _port: int) -> set[str]:
    return {"8.8.8.8"}


def _page(number: int, characters: int, coverage: float = 0.0, garbled: int = 0) -> PageContent:
    return PageContent(
        number=number,
        characters=characters,
        garbled_characters=garbled,
        image_coverage=coverage,
    )


def _scanned_pdf() -> bytes:
    image = Image.new("RGB", (1240, 1754), "white")
    ImageDraw.Draw(image).text((100, 100), "Scanned lecture notes", fill="black")
    stream = io.BytesIO()
    image.save(stream, "PDF", resolution=150)
    return stream.getvalue()


class FakeHostedProcessor(DocumentProcessor):
    id = "fake-hosted"
    label = "Fake OCR"
    description = "Reads scans in tests."
    media_types = frozenset({PDF_MEDIA_TYPE})
    hosted = True
    vendor = "mistral"

    def __init__(self) -> None:
        self.calls = 0

    def version(self) -> str:
        return "1"

    def convert(self, media_type: str, content: bytes) -> Conversion:
        self.calls += 1
        markdown = "# Page 1\n\nRecognised scanned lecture notes about photosynthesis\n"
        return Conversion(markdown=markdown, pages=markdown_pages(markdown))


def _app(tmp_path: Path, fake: FakeHostedProcessor):
    workspace = tmp_path / "course"
    workspace.mkdir(exist_ok=True)

    def with_fake(
        cache_dir: Path, provider_path: Path, store: DocumentProcessingStore
    ) -> DocumentProcessors:
        processors = installed_processors(cache_dir, provider_path, store)
        processors.processors.append(fake)
        return processors

    patcher = patch("course_harness.app.installed_processors", side_effect=with_fake)
    patcher.start()
    app = create_app(
        workspace,
        provider_store_path=tmp_path / "provider",
        library_data_path=tmp_path / "library-data",
        library_cache_path=tmp_path / "library-cache",
        remote_host_resolver=_public_host_resolver,
    )
    return app, patcher


def test_quality_check_flags_scans_screenshots_and_garbled_text() -> None:
    assert check_quality(PDF_MEDIA_TYPE, [_page(1, 2_000), _page(2, 30)]).passed

    scanned = check_quality(PDF_MEDIA_TYPE, [_page(1, 0, 1.0), _page(2, 1_500)])
    assert [finding.kind for finding in scanned.findings] == ["little-text"]
    assert scanned.findings[0].pages == [1]
    assert scanned.findings[0].message.startswith("1 of 2 pages has little or no text")

    slides = check_quality(PPTX_MEDIA_TYPE, [_page(1, 12, 0.9), _page(2, 40)])
    assert [finding.kind for finding in slides.findings] == ["image-heavy"]
    assert "1 of 2 slides is mostly images" in slides.findings[0].message

    garbled = check_quality(PDF_MEDIA_TYPE, [_page(1, 1_000, garbled=100)])
    assert [finding.kind for finding in garbled.findings] == ["garbled-text"]

    # One figure in a long paper is not a reason to reprocess it.
    paper = [_page(number, 3_000) for number in range(1, 20)] + [_page(20, 200, 0.8)]
    assert check_quality(PDF_MEDIA_TYPE, paper).passed


def test_markdown_pages_reads_page_and_slide_sections() -> None:
    pages = markdown_pages("# Page 1\n\nAlpha\n\n# Page 2\n\n# Slide 3\n\nGamma �\n")
    assert [(page.number, page.characters, page.garbled_characters) for page in pages] == [
        (1, 5, 0),
        (2, 0, 0),
        (3, 7, 1),
    ]


@pytest.mark.anyio
async def test_scanned_pdf_is_ready_and_offers_the_preferred_processor(tmp_path: Path) -> None:
    fake = FakeHostedProcessor()
    app, patcher = _app(tmp_path, fake)
    try:
        transport = httpx2.ASGITransport(app=app)
        async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
            quick = await client.post(
                "/api/resources/upload",
                files={"file": ("scan.pdf", _scanned_pdf(), "application/pdf")},
            )
            await client.put(
                "/api/document-processing/preferences",
                json={"preferred_processor": "fake-hosted", "auto_upgrade": False},
            )
            offered = await client.get(f"/api/resources/{quick.json()['resource_id']}")
    finally:
        patcher.stop()

    assert quick.status_code == 201
    assert quick.json()["status"] == "ready"
    assert quick.json()["processing"]["processor"] == "local-quick"
    assert quick.json()["processing"]["quality"]["passed"] is False
    # Nothing better is configured yet, so nothing is suggested.
    assert quick.json()["suggested_processor"] is None
    assert offered.json()["suggested_processor"] == "fake-hosted"
    assert fake.calls == 0


@pytest.mark.anyio
async def test_hosted_reprocessing_needs_consent_and_is_kept_durably(tmp_path: Path) -> None:
    fake = FakeHostedProcessor()
    app, patcher = _app(tmp_path, fake)
    cache_dir = tmp_path / "library-cache"
    try:
        transport = httpx2.ASGITransport(app=app)
        async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
            uploaded = await client.post(
                "/api/resources/upload",
                files={"file": ("scan.pdf", _scanned_pdf(), "application/pdf")},
            )
            resource_id = uploaded.json()["resource_id"]
            refused = await client.post(
                f"/api/resources/{resource_id}/reprocess", json={"processor": "fake-hosted"}
            )
            await client.put("/api/document-processing/consent/mistral")
            hosted = await client.post(
                f"/api/resources/{resource_id}/reprocess", json={"processor": "fake-hosted"}
            )
            assert search_raw(cache_dir, "photosynthesis")
            await client.delete("/api/resources/cache")
            # The Resource remembers its processor, and the vendor is not paid twice.
            restored = await client.post(f"/api/resources/{resource_id}/reprocess")
    finally:
        patcher.stop()

    assert refused.status_code == 409
    assert "Agree to send documents" in refused.json()["detail"]
    assert hosted.status_code == 200
    assert hosted.json()["processing"]["processor"] == "fake-hosted"
    assert hosted.json()["processing"]["hosted"] is True
    assert restored.status_code == 200
    assert restored.json()["processing"]["processor"] == "fake-hosted"
    assert fake.calls == 1
    assert search_raw(cache_dir, "photosynthesis")


@pytest.mark.anyio
async def test_opted_in_upgrade_replaces_a_failed_quick_result(tmp_path: Path) -> None:
    fake = FakeHostedProcessor()
    app, patcher = _app(tmp_path, fake)
    try:
        transport = httpx2.ASGITransport(app=app)
        async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
            await client.put("/api/document-processing/consent/mistral")
            await client.put(
                "/api/document-processing/preferences",
                json={"preferred_processor": "fake-hosted", "auto_upgrade": True},
            )
            passing = await client.post(
                "/api/resources/upload",
                files={"file": ("teaching.pdf", TEACHING_PDF, "application/pdf")},
            )
            scanned = await client.post(
                "/api/resources/upload",
                files={"file": ("scan.pdf", _scanned_pdf(), "application/pdf")},
            )
    finally:
        patcher.stop()

    assert passing.json()["processing"]["processor"] == "local-quick"
    assert scanned.json()["processing"]["processor"] == "fake-hosted"
    assert scanned.json()["suggested_processor"] is None
    assert fake.calls == 1


@pytest.mark.anyio
async def test_a_pinned_representation_is_kept_when_reprocessing(tmp_path: Path) -> None:
    fake = FakeHostedProcessor()
    app, patcher = _app(tmp_path, fake)
    workspace = tmp_path / "course"
    data_dir = tmp_path / "library-data"
    cache_dir = tmp_path / "library-cache"
    try:
        transport = httpx2.ASGITransport(app=app)
        async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
            await client.put("/api/document-processing/consent/mistral")
            pinned = await client.post(
                "/api/resources/upload",
                files={"file": ("teaching.pdf", TEACHING_PDF, "application/pdf")},
            )
            unpinned = await client.post(
                "/api/resources/upload",
                files={"file": ("scan.pdf", _scanned_pdf(), "application/pdf")},
            )
            source = sources_module.admit_source(
                workspace, data_dir, cache_dir, pinned.json()["resource_id"]
            )
            for uploaded in (pinned, unpinned):
                response = await client.post(
                    f"/api/resources/{uploaded.json()['resource_id']}/reprocess",
                    json={"processor": "fake-hosted"},
                )
                assert response.status_code == 200
            content = await client.get(f"/api/sources/{source.id}/content")
            search = await client.post("/api/sources/search", json={"query": "teaching"})
    finally:
        patcher.stop()

    pinned_hash = pinned.json()["snapshot"]["content_hash"]
    unpinned_hash = unpinned.json()["snapshot"]["content_hash"]
    assert source.representation == "local-quick"
    assert "representation: local-quick" in (workspace / "sources.yaml").read_text()
    # The Source keeps reading the representation its Citations point into.
    assert content.text == "# Page 1\n\nPDF teaching evidence\n"
    assert search.json()[0]["source_id"] == source.id
    assert (cache_dir / "derived" / f"{pinned_hash}@local-quick" / "extracted.md").is_file()
    assert "photosynthesis" in (cache_dir / "derived" / pinned_hash / "extracted.md").read_text()
    # Nothing pinned the quick result of the scan, so it was simply replaced.
    assert not (cache_dir / "derived" / f"{unpinned_hash}@local-quick").exists()


def test_sources_without_a_recorded_representation_keep_their_yaml_shape() -> None:
    source = sources_module.Source(
        id="source-000000000001",
        resource_id="resource-000000000001",
        source_version_id="a" * 64,
        label="Notes",
        admitted_at="2026-10-11T00:00:00+00:00",
    )
    serialized = sources_module.serialize_sources_index(
        sources_module.SourcesIndex(sources=[source])
    )
    assert b"representation" not in serialized


@pytest.mark.anyio
async def test_processor_keys_are_stored_but_never_returned(tmp_path: Path) -> None:
    app, patcher = _app(tmp_path, FakeHostedProcessor())
    try:
        transport = httpx2.ASGITransport(app=app)
        async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
            before = await client.get("/api/document-processing")
            saved = await client.put(
                "/api/document-processing/keys/llamaparse", json={"api_key": "llx-secret"}
            )
            unknown = await client.put(
                "/api/document-processing/preferences",
                json={"preferred_processor": "nonexistent", "auto_upgrade": False},
            )
            removed = await client.delete("/api/document-processing/keys/llamaparse")
    finally:
        patcher.stop()

    llamaparse = next(p for p in before.json()["processors"] if p["id"] == "llamaparse")
    assert llamaparse["available"] is False
    assert llamaparse["unavailable_action"] == "add-key"
    assert llamaparse["consented"] is False
    assert "llx-secret" not in saved.text
    assert next(k for k in saved.json()["keys"] if k["vendor"] == "llamaparse")["configured"]
    assert next(p for p in saved.json()["processors"] if p["id"] == "llamaparse")["available"]
    assert unknown.status_code == 422
    assert not next(k for k in removed.json()["keys"] if k["vendor"] == "llamaparse")["configured"]
    assert "llx-secret" not in (tmp_path / "provider" / "document-processing.json").read_text()


def _mock_client(handler):
    def factory(*, timeout: float, follow_redirects: bool = False) -> httpx2.AsyncClient:
        return httpx2.AsyncClient(transport=httpx2.MockTransport(handler), timeout=timeout)

    return patch.object(hosted_processors, "product_connector_client", side_effect=factory)


def test_llamaparse_uploads_polls_and_sections_pages() -> None:
    requests: list[httpx2.Request] = []

    def handler(request: httpx2.Request) -> httpx2.Response:
        requests.append(request)
        assert request.headers["Authorization"] == "Bearer llx-key"
        if request.method == "POST":
            return httpx2.Response(200, json={"id": "job-1", "status": "PENDING"})
        if len(requests) == 2:
            return httpx2.Response(200, json={"job": {"id": "job-1", "status": "RUNNING"}})
        return httpx2.Response(
            200,
            json={
                "job": {"id": "job-1", "status": "COMPLETED"},
                "markdown": {
                    "pages": [
                        {"page_number": 2, "markdown": "Second", "success": True},
                        {"page_number": 1, "markdown": "| a | b |", "success": True},
                    ]
                },
            },
        )

    with _mock_client(handler), patch.object(hosted_processors.asyncio, "sleep"):
        conversion = LlamaParseProcessor("llx-key").convert(PPTX_MEDIA_TYPE, b"pptx")

    assert requests[0].url.path == "/api/v2/parse/upload"
    assert b'"tier": "cost_effective"' in requests[0].content
    assert requests[-1].url.params["expand"] == "markdown"
    assert conversion.markdown == "# Slide 1\n\n| a | b |\n\n# Slide 2\n\nSecond\n"
    assert [page.number for page in conversion.pages] == [1, 2]


def test_mistral_ocr_sends_the_pdf_inline_and_reports_a_bad_key() -> None:
    sent: list[dict] = []

    def handler(request: httpx2.Request) -> httpx2.Response:
        sent.append(json.loads(request.content))
        return httpx2.Response(
            200,
            json={
                "model": "mistral-ocr-latest",
                "pages": [{"index": 0, "markdown": "Recognised text"}],
                "usage_info": {"pages_processed": 1},
            },
        )

    with _mock_client(handler):
        conversion = MistralOcrProcessor("mistral-key").convert(PDF_MEDIA_TYPE, b"%PDF-1.4")

    assert sent[0]["model"] == "mistral-ocr-latest"
    assert sent[0]["document"]["document_url"].startswith("data:application/pdf;base64,")
    assert conversion.markdown == "# Page 1\n\nRecognised text\n"

    def rejected(_request: httpx2.Request) -> httpx2.Response:
        return httpx2.Response(401, json={"message": "Unauthorized"})

    with _mock_client(rejected), pytest.raises(RuntimeError, match="rejected the API key"):
        MistralOcrProcessor("bad").convert(PDF_MEDIA_TYPE, b"%PDF-1.4")


def test_pptx_screenshot_slides_fail_the_quality_check(tmp_path: Path) -> None:
    from course_harness.docling_parser import convert_document  # noqa: PLC0415

    presentation = PPTXPresentation()
    image = io.BytesIO()
    Image.new("RGB", (640, 480), "white").save(image, "PNG")
    image.seek(0)
    screenshot = presentation.slides.add_slide(presentation.slide_layouts[6])
    screenshot.shapes.add_picture(
        image, 0, 0, width=presentation.slide_width, height=presentation.slide_height
    )
    titled = presentation.slides.add_slide(presentation.slide_layouts[5])
    titled.shapes.title.text = "A slide with text"
    stream = io.BytesIO()
    presentation.save(stream)

    conversion = convert_document(PPTX_MEDIA_TYPE, stream.getvalue(), tmp_path)
    report = check_quality(PPTX_MEDIA_TYPE, conversion.pages)

    assert [finding.kind for finding in report.findings] == ["image-heavy"]
    assert report.findings[0].pages == [1]
