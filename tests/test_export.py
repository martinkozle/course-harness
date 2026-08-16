from __future__ import annotations

from io import BytesIO
from pathlib import Path

import httpx2
import pytest
from pptx import Presentation as PPTXPresentation

from course_harness.app import create_app
from course_harness.course_plan import (
    CoursePlanInput,
    LectureInput,
    create_course_plan,
    create_course_plan_file,
    initialize_workspace_history,
)
from course_harness.export import export_presentation, validate_export_mapping
from course_harness.presentation import BulletsSlide, Presentation
from course_harness.template_profiles import (
    TemplateLayoutMapping,
    TemplateProfile,
    _generate_calibration_deck,
)


def _build_presentation(workspace: Path, lecture_id: str, slides: list[dict]):
    from course_harness.presentation import SLIDE_CLASSES_BY_LAYOUT, write_presentation
    from course_harness.presentation import Presentation as PresModel

    hydrated = []
    for s in slides:
        layout = s.get("layout", "title")
        cls = SLIDE_CLASSES_BY_LAYOUT[layout]
        hydrated.append(cls(**s))
    pres = PresModel(
        id="presentation-abc123def456",
        lecture_id=lecture_id,
        slides=hydrated,
    )
    write_presentation(workspace, pres)
    return pres


def test_export_uses_corrected_template_slots(tmp_path: Path) -> None:
    template_path = tmp_path / "template.pptx"
    PPTXPresentation().save(str(template_path))
    profile = TemplateProfile(
        id="tpl-000000000001",
        name="Corrected slots",
        version=1,
        template_filename="template.pptx",
        slide_width=12_192_000,
        slide_height=6_858_000,
        slide_count=11,
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="bullets",
                template_layout_index=3,
                confidence=1,
                rationale="Course Author correction",
                slot_mappings={"title": 0, "body": 2},
            )
        ],
    )
    presentation = Presentation(
        id="presentation-abc123def456",
        lecture_id="lecture-abc123def456",
        slides=[
            BulletsSlide(
                id="slide-abc123def456",
                title="Mapped title",
                bullets=["Mapped body"],
            )
        ],
    )

    exported = PPTXPresentation(BytesIO(export_presentation(presentation, profile, template_path)))

    placeholders = {
        shape.placeholder_format.idx: shape.text_frame.text
        for shape in exported.slides[0].placeholders
        if shape.has_text_frame
    }
    assert placeholders[0] == "Mapped title"
    assert placeholders[1] == ""
    assert placeholders[2] == "Mapped body"


def test_calibration_deck_uses_corrected_template_slots(tmp_path: Path) -> None:
    template_path = tmp_path / "template.pptx"
    PPTXPresentation().save(str(template_path))
    profile = TemplateProfile(
        id="tpl-000000000001",
        name="Corrected slots",
        version=1,
        template_filename="template.pptx",
        slide_width=12_192_000,
        slide_height=6_858_000,
        slide_count=11,
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="bullets",
                template_layout_index=3,
                confidence=1,
                rationale="Course Author correction",
                slot_mappings={"title": 0, "body": 2},
            )
        ],
    )

    calibrated = PPTXPresentation(BytesIO(_generate_calibration_deck(profile, template_path)))

    placeholders = {
        shape.placeholder_format.idx: shape.text_frame.text
        for shape in calibrated.slides[-1].placeholders
        if shape.has_text_frame
    }
    assert placeholders[0] == "[bullets] Sample Title"
    assert placeholders[1] == ""
    assert "Sample body text" in placeholders[2]


def test_calibration_deck_uses_export_fallback_for_unmapped_slots(tmp_path: Path) -> None:
    template_path = tmp_path / "template.pptx"
    PPTXPresentation().save(str(template_path))
    profile = TemplateProfile(
        id="tpl-000000000001",
        name="Partial slots",
        version=1,
        template_filename="template.pptx",
        slide_width=12_192_000,
        slide_height=6_858_000,
        slide_count=11,
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="bullets",
                template_layout_index=1,
                confidence=1,
                rationale="Keep the inferred body fallback",
                slot_mappings={"title": 0},
            )
        ],
    )

    calibrated = PPTXPresentation(BytesIO(_generate_calibration_deck(profile, template_path)))

    placeholders = {
        shape.placeholder_format.idx: shape.text_frame.text
        for shape in calibrated.slides[-1].placeholders
        if shape.has_text_frame
    }
    assert placeholders[0] == "[bullets] Sample Title"
    assert "Sample body text" in placeholders[1]


def test_calibration_deck_uses_two_column_fallback_for_unmapped_columns(tmp_path: Path) -> None:
    template_path = tmp_path / "template.pptx"
    PPTXPresentation().save(str(template_path))
    profile = TemplateProfile(
        id="tpl-000000000001",
        name="Partial columns",
        version=1,
        template_filename="template.pptx",
        slide_width=12_192_000,
        slide_height=6_858_000,
        slide_count=11,
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="two_column",
                template_layout_index=3,
                confidence=1,
                rationale="Keep inferred column fallbacks",
                slot_mappings={"title": 0},
            )
        ],
    )

    calibrated = PPTXPresentation(BytesIO(_generate_calibration_deck(profile, template_path)))

    placeholders = {
        shape.placeholder_format.idx: shape.text_frame.text
        for shape in calibrated.slides[-1].placeholders
        if shape.has_text_frame
    }
    assert placeholders[1] == "Sample left-column content."
    assert placeholders[2] == "Sample right-column content."


def test_validation_blocks_incomplete_two_column_slot_mapping(tmp_path: Path) -> None:
    template_path = tmp_path / "template.pptx"
    PPTXPresentation().save(str(template_path))
    profile = TemplateProfile(
        id="tpl-000000000001",
        name="Incomplete columns",
        version=1,
        template_filename="template.pptx",
        slide_width=12_192_000,
        slide_height=6_858_000,
        slide_count=11,
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="two_column",
                template_layout_index=3,
                confidence=1,
                rationale="Incomplete correction",
                slot_mappings={"title": 0, "left": 1},
            )
        ],
    )

    issues = validate_export_mapping(profile, template_path)

    assert issues == [
        {
            "level": "blocking",
            "message": (
                "'two_column' → 'Two Content' (index 3): "
                "both left and right slots must be configured"
            ),
        }
    ]


def test_validation_blocks_two_column_layout_with_one_horizontal_group(tmp_path: Path) -> None:
    template_path = tmp_path / "template.pptx"
    template = PPTXPresentation()
    layout = template.slide_layouts[3]
    layout.placeholders[2].left = layout.placeholders[1].left
    template.save(str(template_path))
    profile = TemplateProfile(
        id="tpl-000000000001",
        name="Stacked content",
        version=1,
        template_filename="template.pptx",
        slide_width=12_192_000,
        slide_height=6_858_000,
        slide_count=11,
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="two_column",
                template_layout_index=3,
                confidence=1,
                rationale="Not actually two columns",
            )
        ],
    )

    issues = validate_export_mapping(profile, template_path)

    assert {
        "level": "blocking",
        "message": (
            "'two_column' → 'Two Content' (index 3): only 1 content column group(s) "
            "found, two_column needs at least 2"
        ),
    } in issues


def test_validation_blocks_missing_required_content_placeholder(tmp_path: Path) -> None:
    template_path = tmp_path / "template.pptx"
    PPTXPresentation().save(str(template_path))
    profile = TemplateProfile(
        id="tpl-000000000001",
        name="Missing content",
        version=1,
        template_filename="template.pptx",
        slide_width=12_192_000,
        slide_height=6_858_000,
        slide_count=11,
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="bullets",
                template_layout_index=5,
                confidence=1,
                rationale="Unusable title-only layout",
            )
        ],
    )

    issues = validate_export_mapping(profile, template_path)

    assert {
        "level": "blocking",
        "message": (
            "'bullets' → 'Title Only' (index 5): slot 'body' has no compatible "
            "placeholder (expected BODY or OBJECT)"
        ),
    } in issues


def test_validation_uses_type_fallback_when_content_index_is_not_one(tmp_path: Path) -> None:
    template_path = tmp_path / "template.pptx"
    template = PPTXPresentation()
    content_placeholder = template.slide_layouts[1].placeholders[1]
    content_placeholder.placeholder_format._ph.idx = 2
    template.save(str(template_path))
    profile = TemplateProfile(
        id="tpl-000000000001",
        name="Nonstandard content index",
        version=1,
        template_filename="template.pptx",
        slide_width=12_192_000,
        slide_height=6_858_000,
        slide_count=11,
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="bullets",
                template_layout_index=1,
                confidence=1,
                rationale="Usable type-based fallback",
            )
        ],
    )

    issues = validate_export_mapping(profile, template_path)

    assert issues == []


def test_validation_blocks_slot_collisions_and_unsuitable_column_types(tmp_path: Path) -> None:
    template_path = tmp_path / "template.pptx"
    PPTXPresentation().save(str(template_path))
    profile = TemplateProfile(
        id="tpl-000000000001",
        name="Colliding columns",
        version=1,
        template_filename="template.pptx",
        slide_width=12_192_000,
        slide_height=6_858_000,
        slide_count=11,
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="two_column",
                template_layout_index=3,
                confidence=1,
                rationale="Invalid correction",
                slot_mappings={"title": 0, "left": 0, "right": 2},
            )
        ],
    )

    issues = validate_export_mapping(profile, template_path)

    blocking_messages = [issue["message"] for issue in issues if issue["level"] == "blocking"]
    assert any(
        "slots 'left' and 'title' use the same placeholder idx=0" in message
        for message in blocking_messages
    )
    assert any(
        "slot 'left' maps to type TITLE, expected BODY or OBJECT" in message
        for message in blocking_messages
    )


@pytest.mark.anyio
async def test_export_title_slide(tmp_path: Path) -> None:
    workspace = tmp_path / "export-course"
    workspace.mkdir()
    plan = create_course_plan(
        CoursePlanInput(
            title="Export Test",
            audience="Test",
            lectures=[LectureInput(title="L1")],
        )
    )
    initialize_workspace_history(workspace)
    create_course_plan_file(workspace, plan)
    _build_presentation(
        workspace,
        plan.lectures[0].id,
        [{"layout": "title", "id": "slide-000000000001", "title": "Hello", "subtitle": "World"}],
    )

    transport = httpx2.ASGITransport(app=create_app(workspace))
    async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.get(f"/api/presentations/{plan.lectures[0].id}/export")

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )

    pptx = PPTXPresentation(BytesIO(response.content))
    assert len(pptx.slides) == 1
    slide = pptx.slides[0]
    assert slide.slide_layout.name is not None
    text_runs = [
        run.text
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    ]
    assert "Hello" in text_runs
    assert "World" in text_runs


@pytest.mark.anyio
async def test_export_all_layouts(tmp_path: Path) -> None:
    workspace = tmp_path / "export-all"
    workspace.mkdir()
    plan = create_course_plan(
        CoursePlanInput(
            title="All Layouts",
            audience="Test",
            lectures=[LectureInput(title="L1")],
        )
    )
    initialize_workspace_history(workspace)
    create_course_plan_file(workspace, plan)
    lid = plan.lectures[0].id
    all_slides = [
        {
            "layout": "title",
            "id": "slide-00000000000a",
            "title": "Title Slide",
            "subtitle": "Subtitle",
        },
        {"layout": "section", "id": "slide-00000000000b", "title": "Section Header"},
        {
            "layout": "bullets",
            "id": "slide-00000000000c",
            "title": "Bullets",
            "bullets": ["One", "Two", "Three"],
        },
        {
            "layout": "two_column",
            "id": "slide-00000000000d",
            "title": "Two Column",
            "left_content": "Left side",
            "right_content": "Right side",
        },
        {
            "layout": "big_statement",
            "id": "slide-00000000000e",
            "title": "Minor header",
            "statement": "BIG STATEMENT",
        },
        {"layout": "closing", "id": "slide-00000000000f", "title": "Closing", "text": "Thanks"},
        {
            "layout": "code",
            "id": "slide-000000000010",
            "title": "Code",
            "code": "def hello():\n    pass",
            "language": "python",
        },
        {
            "layout": "quote",
            "id": "slide-000000000011",
            "title": "Quote",
            "quote": "To be or not",
            "attribution": "Bard",
        },
    ]
    _build_presentation(workspace, lid, all_slides)

    transport = httpx2.ASGITransport(app=create_app(workspace))
    async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.get(f"/api/presentations/{lid}/export")

    assert response.status_code == 200
    pptx = PPTXPresentation(BytesIO(response.content))
    assert len(pptx.slides) == 8
    text_runs = [
        run.text
        for slide in pptx.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    ]
    combined = " ".join(text_runs)
    assert "Title Slide" in combined
    assert "BIG STATEMENT" in combined
    assert "Left side" in combined
    assert "To be or not" in combined
    assert "Bard" in combined
    assert "def hello():" in combined


@pytest.mark.anyio
async def test_export_citations_and_notes(tmp_path: Path) -> None:
    workspace = tmp_path / "export-cite"
    workspace.mkdir()
    plan = create_course_plan(
        CoursePlanInput(
            title="Citations Test",
            audience="Test",
            lectures=[LectureInput(title="L1")],
        )
    )
    initialize_workspace_history(workspace)
    create_course_plan_file(workspace, plan)
    _build_presentation(
        workspace,
        plan.lectures[0].id,
        [
            {
                "layout": "title",
                "id": "slide-000000000001",
                "title": "Slide 1",
                "speaker_notes": "Speaker note content here.",
                "citations": [
                    {
                        "source_id": "src-aaaa1111bbbb",
                        "label": "Paper A",
                        "url": "https://example.com/a",
                    },
                    {
                        "source_id": "src-cccc2222dddd",
                        "label": "Paper B",
                        "line_start": 10,
                        "line_end": 20,
                    },
                ],
            },
        ],
    )

    transport = httpx2.ASGITransport(app=create_app(workspace))
    async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.get(f"/api/presentations/{plan.lectures[0].id}/export")

    assert response.status_code == 200
    pptx = PPTXPresentation(BytesIO(response.content))
    slide = pptx.slides[0]
    has_notes = False
    try:
        notes = slide.notes_slide
        notes_text = notes.notes_text_frame.text
        has_notes = True
    except Exception:
        notes_text = ""
    if has_notes:
        assert "Speaker note content here" in notes_text
        assert "Paper A" in notes_text
        assert "lines 11" in notes_text
    shape_texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    shape_texts.append(run.text)
    all_text = " ".join(shape_texts)
    assert "Paper A" in all_text
    assert all_text.count("Paper A") == 1


@pytest.mark.anyio
async def test_export_uses_course_pinned_profile_version(tmp_path: Path) -> None:
    from course_harness.template_profiles import (
        TemplateLayoutMapping,
        TemplateProfile,
        profile_dir,
        write_profile_version,
    )

    workspace = tmp_path / "export-pinned"
    workspace.mkdir()
    plan = create_course_plan(
        CoursePlanInput(
            title="Pinned export",
            audience="Test",
            lectures=[LectureInput(title="L1")],
        )
    ).model_copy(
        update={
            "template_profile_id": "tpl-000000000001",
            "template_profile_version": 1,
        }
    )
    initialize_workspace_history(workspace)
    create_course_plan_file(workspace, plan)
    _build_presentation(
        workspace,
        plan.lectures[0].id,
        [{"layout": "title", "id": "slide-000000000001", "title": "Pinned"}],
    )

    templates_data = tmp_path / "templates"
    v1 = TemplateProfile(
        id="tpl-000000000001",
        name="Pinned template",
        template_filename="template.pptx",
        slide_width=12192000,
        slide_height=6858000,
        slide_count=11,
        version=1,
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="title",
                template_layout_index=0,
                confidence=1,
                rationale="Version one",
            )
        ],
    )
    v2 = TemplateProfile(
        id="tpl-000000000001",
        name="Pinned template",
        template_filename="template.pptx",
        slide_width=12192000,
        slide_height=6858000,
        slide_count=11,
        version=2,
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="title",
                template_layout_index=2,
                confidence=1,
                rationale="Version two",
            )
        ],
    )
    write_profile_version(templates_data, v1)
    write_profile_version(templates_data, v2)
    template_file = profile_dir(templates_data, v1.id) / "template.pptx"
    PPTXPresentation().save(str(template_file))

    app = create_app(
        workspace,
        templates_data_path=templates_data,
        templates_cache_path=tmp_path / "template-cache",
    )
    transport = httpx2.ASGITransport(app=app)
    async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.get(f"/api/presentations/{plan.lectures[0].id}/export")

    assert response.status_code == 200
    exported = PPTXPresentation(BytesIO(response.content))
    assert exported.slides[0].slide_layout.name == "Title Slide"


@pytest.mark.anyio
async def test_export_slide_order(tmp_path: Path) -> None:
    workspace = tmp_path / "export-order"
    workspace.mkdir()
    plan = create_course_plan(
        CoursePlanInput(
            title="Order",
            audience="Test",
            lectures=[LectureInput(title="L1")],
        )
    )
    initialize_workspace_history(workspace)
    create_course_plan_file(workspace, plan)
    _build_presentation(
        workspace,
        plan.lectures[0].id,
        [
            {"layout": "title", "id": "slide-000000000001", "title": "First"},
            {"layout": "title", "id": "slide-000000000002", "title": "Second"},
            {"layout": "title", "id": "slide-000000000003", "title": "Third"},
        ],
    )

    transport = httpx2.ASGITransport(app=create_app(workspace))
    async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.get(f"/api/presentations/{plan.lectures[0].id}/export")

    assert response.status_code == 200
    pptx = PPTXPresentation(BytesIO(response.content))
    slide_titles = []
    for s in pptx.slides:
        for shape in s.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    slide_titles.append(txt)
    assert slide_titles == ["First", "Second", "Third"]


@pytest.mark.anyio
async def test_export_skips_archived(tmp_path: Path) -> None:
    workspace = tmp_path / "export-archived"
    workspace.mkdir()
    plan = create_course_plan(
        CoursePlanInput(
            title="Archived",
            audience="Test",
            lectures=[LectureInput(title="L1")],
        )
    )
    initialize_workspace_history(workspace)
    create_course_plan_file(workspace, plan)
    _build_presentation(
        workspace,
        plan.lectures[0].id,
        [
            {"layout": "title", "id": "slide-000000000001", "title": "Active"},
            {"layout": "title", "id": "slide-000000000002", "title": "Hidden", "archived": True},
        ],
    )

    transport = httpx2.ASGITransport(app=create_app(workspace))
    async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.get(f"/api/presentations/{plan.lectures[0].id}/export")

    assert response.status_code == 200
    pptx = PPTXPresentation(BytesIO(response.content))
    assert len(pptx.slides) == 1
    titles = []
    for s in pptx.slides:
        for shape in s.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    titles.append(txt)
    assert titles == ["Active"]


@pytest.mark.anyio
async def test_export_no_presentation_returns_404(tmp_path: Path) -> None:
    workspace = tmp_path / "export-404"
    workspace.mkdir()
    plan = create_course_plan(
        CoursePlanInput(
            title="No Pres",
            audience="Test",
            lectures=[LectureInput(title="L1")],
        )
    )
    initialize_workspace_history(workspace)
    create_course_plan_file(workspace, plan)

    transport = httpx2.ASGITransport(app=create_app(workspace))
    async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.get(f"/api/presentations/{plan.lectures[0].id}/export")

    assert response.status_code == 404


@pytest.mark.anyio
async def test_export_empty_presentation(tmp_path: Path) -> None:
    workspace = tmp_path / "export-empty"
    workspace.mkdir()
    plan = create_course_plan(
        CoursePlanInput(
            title="Empty",
            audience="Test",
            lectures=[LectureInput(title="L1")],
        )
    )
    initialize_workspace_history(workspace)
    create_course_plan_file(workspace, plan)
    _build_presentation(workspace, plan.lectures[0].id, [])

    transport = httpx2.ASGITransport(app=create_app(workspace))
    async with httpx2.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.get(f"/api/presentations/{plan.lectures[0].id}/export")

    assert response.status_code == 200
    pptx = PPTXPresentation(BytesIO(response.content))
    assert len(pptx.slides) == 0
