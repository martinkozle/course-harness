from pathlib import Path

import pytest
from pptx import Presentation as PPTXPresentation

from course_harness.template_inspect import _name_score, inspect_template, map_semantic_layouts


def _make_test_template() -> bytes:
    prs = PPTXPresentation()
    # Verify we have standard layouts
    assert len(prs.slide_layouts) >= 9
    buffer = __import__("io").BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def test_template_path(tmp_path: Path) -> Path:
    template_path = tmp_path / "test.pptx"
    template_path.write_bytes(_make_test_template())
    return template_path


def test_inspect_template_returns_layout_count(test_template_path: Path) -> None:
    result = inspect_template(test_template_path)
    layouts = result["layouts"]
    assert len(layouts) >= 9
    for layout in layouts:
        assert "index" in layout
        assert "name" in layout
        assert isinstance(layout["index"], int)
        assert isinstance(layout["name"], str)


def test_inspect_template_returns_dimensions(test_template_path: Path) -> None:
    result = inspect_template(test_template_path)
    assert result["slide_width"] > 0
    assert result["slide_height"] > 0


def test_inspect_template_placeholders_have_correct_types(test_template_path: Path) -> None:
    result = inspect_template(test_template_path)
    title_layout = result["layouts"][0]
    ph_types = {ph["type"] for ph in title_layout["placeholders"]}
    # CENTER_TITLE = 3 and SUBTITLE = 4 are expected on a title slide
    assert {3, 4}.issubset(ph_types)


def test_map_semantic_layouts_covers_all_semantics(test_template_path: Path) -> None:
    result = map_semantic_layouts(inspect_template(test_template_path))
    expected = {
        "title",
        "section",
        "bullets",
        "two_column",
        "big_statement",
        "closing",
        "code",
        "image",
        "quote",
    }
    actual = {m["semantic_layout"] for m in result}
    assert actual == expected


def test_map_semantic_layouts_all_have_confidence_and_rationale(test_template_path: Path) -> None:
    result = map_semantic_layouts(inspect_template(test_template_path))
    for m in result:
        assert 0.0 <= m["confidence"] <= 1.0
        assert len(m["rationale"]) > 0
        assert isinstance(m["template_layout_index"], int)
        assert m["template_layout_index"] >= 0


def test_map_semantic_layouts_title_gets_index_zero(test_template_path: Path) -> None:
    result = map_semantic_layouts(inspect_template(test_template_path))
    title_mapping = next(m for m in result if m["semantic_layout"] == "title")
    assert title_mapping["template_layout_index"] == 0
    assert title_mapping["confidence"] >= 0.8


def test_name_score_exact_match() -> None:
    assert _name_score("Title Slide", "title") >= 0.8


def test_name_score_partial_match() -> None:
    assert _name_score("Section Header", "section") >= 0.5


def test_name_score_no_match() -> None:
    assert _name_score("Title and Content", "image") <= 0.5
