import hashlib
import json
import subprocess
from pathlib import Path

import pytest
from pptx import Presentation as PPTXPresentation
from pydantic import ValidationError

from course_harness.course_plan import CoursePlan, Lecture, write_course_plan
from course_harness.presentation import Presentation, TitleSlide, write_presentation
from course_harness.release_service import (
    PublishReleaseRequest,
    ReleaseConflict,
    ReleaseError,
    list_releases,
    publish_release,
    read_release,
    read_release_artifact,
    regenerate_release_artifact,
)
from course_harness.release_validation import ReleaseSelection, Waiver
from course_harness.sources import Source, SourcesIndex, write_sources_index
from course_harness.template_profiles import (
    BUILTIN_DEFAULT_ID,
    TemplateLayoutMapping,
    TemplateProfile,
    profile_dir,
    write_profile_version,
)
from course_harness.workspace_history import (
    RevisionCreateRequest,
    create_revision,
    record_app_authored_state,
)

LECTURE_ID = "lecture-aaaaaaaaaaaa"
PRESENTATION_ID = "presentation-aaaaaaaaaaaa"
SECOND_LECTURE_ID = "lecture-bbbbbbbbbbbb"


def _clean_course(
    workspace: Path,
    *,
    second_lecture: bool = False,
    template_profile: TemplateProfile | None = None,
    sources: SourcesIndex | None = None,
) -> None:
    workspace.mkdir()
    subprocess.run(["git", "init", "--quiet", "--initial-branch=main", str(workspace)], check=True)
    write_course_plan(
        workspace,
        CoursePlan(
            id="course-aaaaaaaaaaaa",
            title="Causal inference",
            audience="Graduate students",
            template_profile_id=template_profile.id if template_profile else None,
            template_profile_version=template_profile.version if template_profile else None,
            lectures=[
                Lecture(
                    id=LECTURE_ID,
                    title="Foundations",
                    presentation_id=PRESENTATION_ID,
                ),
                *([Lecture(id=SECOND_LECTURE_ID, title="Applications")] if second_lecture else []),
            ],
        ),
    )
    write_presentation(
        workspace,
        Presentation(
            id=PRESENTATION_ID,
            lecture_id=LECTURE_ID,
            slides=[TitleSlide(id="slide-aaaaaaaaaaaa", title="Foundations")],
        ),
    )
    if sources is not None:
        write_sources_index(workspace, sources)
    create_revision(workspace, RevisionCreateRequest(summary="Initial Course"))
    record_app_authored_state(workspace)


def _custom_template(templates: Path) -> TemplateProfile:
    profile = TemplateProfile(
        id="tpl-aaaaaaaaaaaa",
        name="Institutional",
        version=3,
        template_filename="institutional.pptx",
        slide_width=12_192_000,
        slide_height=6_858_000,
        slide_count=11,
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="title",
                template_layout_index=0,
                confidence=1,
                rationale="Course Author mapping",
            )
        ],
    )
    write_profile_version(templates, profile)
    PPTXPresentation().save(str(profile_dir(templates, profile.id) / "template.pptx"))
    return profile


def _replace_release_tag(workspace: Path, slug: str, commit_oid: str, message: bytes) -> None:
    subprocess.run(
        ["git", "tag", "--delete", f"course-release/{slug}"],
        cwd=workspace,
        check=True,
        capture_output=True,
    )
    subprocess.run(
        [
            "git",
            "-c",
            "tag.gpgSign=false",
            "-c",
            "user.name=Course Harness tests",
            "-c",
            "user.email=course-harness-tests@example.invalid",
            "tag",
            "--annotate",
            "--cleanup=verbatim",
            "--file=-",
            f"course-release/{slug}",
            commit_oid,
        ],
        cwd=workspace,
        input=message,
        check=True,
    )


def test_publish_release_exports_artifacts_and_records_a_bounded_annotated_tag(
    tmp_path: Path,
) -> None:
    workspace = tmp_path / "course"
    release_data = tmp_path / "releases"
    _clean_course(workspace, second_lecture=True)

    release = publish_release(
        workspace=workspace,
        release_data_root=release_data,
        templates_data_root=tmp_path / "templates",
        request=PublishReleaseRequest(
            slug="fall-2026",
            name="Fall 2026",
            selection=ReleaseSelection(lecture_ids=[LECTURE_ID], artifact_ids=[PRESENTATION_ID]),
        ),
    )

    artifact = (
        release_data
        / "course-aaaaaaaaaaaa"
        / "fall-2026"
        / "artifacts"
        / (f"{PRESENTATION_ID}.pptx")
    )
    assert artifact.read_bytes()[:2] == b"PK"
    assert release.tag == "course-release/fall-2026"
    commit_oid = subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=workspace,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    assert release.revision_id == f"revision-{commit_oid}"
    assert release.commit_oid == commit_oid
    assert release.template_profile.id == BUILTIN_DEFAULT_ID
    assert release.template_profile.version == 1
    assert release.included_lecture_ids == [LECTURE_ID]
    assert release.planned_unpublished_lecture_ids == [SECOND_LECTURE_ID]
    assert release.artifacts[0].sha256 == hashlib.sha256(artifact.read_bytes()).hexdigest()
    assert list_releases(workspace) == [release]
    assert read_release(workspace, "fall-2026") == release

    tag_type = subprocess.run(
        ["git", "cat-file", "-t", "refs/tags/course-release/fall-2026"],
        cwd=workspace,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    assert tag_type == "tag"
    assert (
        subprocess.run(
            ["git", "rev-parse", "course-release/fall-2026^{commit}"],
            cwd=workspace,
            check=True,
            capture_output=True,
            text=True,
        ).stdout.strip()
        == commit_oid
    )

    (workspace / "course.yaml").write_text("not: [valid", encoding="utf-8")
    assert read_release(workspace, "fall-2026") == release
    assert list_releases(workspace) == [release]
    stored = read_release_artifact(
        workspace=workspace,
        release_data_root=release_data,
        slug="fall-2026",
        artifact_id=PRESENTATION_ID,
    )
    assert (
        regenerate_release_artifact(
            workspace=workspace,
            release_data_root=release_data,
            slug="fall-2026",
            artifact_id=PRESENTATION_ID,
        )
        == stored
    )


def test_publish_release_never_overwrites_a_name_or_its_artifacts(tmp_path: Path) -> None:
    workspace = tmp_path / "course"
    release_data = tmp_path / "releases"
    _clean_course(workspace)
    request = PublishReleaseRequest(
        slug="fall-2026",
        name="Fall 2026",
        selection=ReleaseSelection(lecture_ids=[LECTURE_ID], artifact_ids=[]),
    )
    publish_release(
        workspace=workspace,
        release_data_root=release_data,
        templates_data_root=tmp_path / "templates",
        request=request,
    )

    with pytest.raises(ReleaseConflict, match="already exists"):
        publish_release(
            workspace=workspace,
            release_data_root=release_data,
            templates_data_root=tmp_path / "templates",
            request=request,
        )


def test_publish_release_rejects_workspace_drift_and_storage_inside_workspace(
    tmp_path: Path,
) -> None:
    workspace = tmp_path / "course"
    _clean_course(workspace)
    request = PublishReleaseRequest(
        slug="fall-2026",
        name="Fall 2026",
        selection=ReleaseSelection(lecture_ids=[LECTURE_ID]),
    )
    original = (workspace / "course.yaml").read_text(encoding="utf-8")
    (workspace / "course.yaml").write_text(original.replace("Causal", "Applied"), encoding="utf-8")

    with pytest.raises(ReleaseError, match="provenance-clean"):
        publish_release(
            workspace=workspace,
            release_data_root=tmp_path / "releases",
            templates_data_root=tmp_path / "templates",
            request=request,
        )

    (workspace / "course.yaml").write_text(original, encoding="utf-8")
    with pytest.raises(ReleaseError, match="outside"):
        publish_release(
            workspace=workspace,
            release_data_root=workspace / ".release-data",
            templates_data_root=tmp_path / "templates",
            request=request,
        )


def test_release_identifiers_and_tag_schema_are_rejected_safely(tmp_path: Path) -> None:
    with pytest.raises(ValidationError):
        PublishReleaseRequest(
            slug="../unsafe",
            name="Unsafe",
            selection=ReleaseSelection(lecture_ids=[LECTURE_ID]),
        )

    workspace = tmp_path / "course"
    _clean_course(workspace)
    subprocess.run(
        ["git", "-c", "tag.gpgSign=false", "tag", "course-release/not-annotated"],
        cwd=workspace,
        check=True,
    )
    with pytest.raises(ReleaseError, match="Git operation failed"):
        read_release(workspace, "not-annotated")


def test_publish_release_translates_a_stale_waiver_to_a_release_error(tmp_path: Path) -> None:
    workspace = tmp_path / "course"
    _clean_course(workspace)

    with pytest.raises(ReleaseError, match="does not match a current warning"):
        publish_release(
            workspace=workspace,
            release_data_root=tmp_path / "releases",
            templates_data_root=tmp_path / "templates",
            request=PublishReleaseRequest(
                slug="waived",
                name="Waived",
                selection=ReleaseSelection(lecture_ids=[LECTURE_ID]),
                waivers=[Waiver(finding_id="stale", justification="Accepted")],
            ),
        )


def test_release_artifact_access_rejects_changed_bytes_and_symlinks(tmp_path: Path) -> None:
    workspace = tmp_path / "course"
    release_data = tmp_path / "releases"
    _clean_course(workspace)
    publish_release(
        workspace=workspace,
        release_data_root=release_data,
        templates_data_root=tmp_path / "templates",
        request=PublishReleaseRequest(
            slug="fall-2026",
            name="Fall 2026",
            selection=ReleaseSelection(lecture_ids=[LECTURE_ID], artifact_ids=[PRESENTATION_ID]),
        ),
    )
    artifact = (
        release_data / "course-aaaaaaaaaaaa" / "fall-2026" / "artifacts" / f"{PRESENTATION_ID}.pptx"
    )
    artifact.write_bytes(b"changed")
    with pytest.raises(ReleaseError, match="does not match"):
        read_release_artifact(
            workspace=workspace,
            release_data_root=release_data,
            slug="fall-2026",
            artifact_id=PRESENTATION_ID,
        )
    artifact.unlink()
    artifact.symlink_to(workspace / "course.yaml")
    with pytest.raises(ReleaseError, match="cannot be read safely"):
        read_release_artifact(
            workspace=workspace,
            release_data_root=release_data,
            slug="fall-2026",
            artifact_id=PRESENTATION_ID,
        )


def test_regeneration_uses_tagged_state_and_stored_custom_template_only(tmp_path: Path) -> None:
    workspace = tmp_path / "course"
    release_data = tmp_path / "releases"
    templates = tmp_path / "templates"
    profile = _custom_template(templates)
    sources = SourcesIndex(
        sources=[
            Source(
                id="source-aaaaaaaaaaaa",
                resource_id="resource-aaaaaaaaaaaa",
                source_version_id="a" * 64,
                label="Foundations",
                admitted_at="2026-09-15T00:00:00+00:00",
            )
        ]
    )
    _clean_course(workspace, template_profile=profile, sources=sources)
    release = publish_release(
        workspace=workspace,
        release_data_root=release_data,
        templates_data_root=templates,
        request=PublishReleaseRequest(
            slug="institutional",
            name="Institutional",
            selection=ReleaseSelection(lecture_ids=[LECTURE_ID], artifact_ids=[PRESENTATION_ID]),
        ),
    )
    original = read_release_artifact(
        workspace=workspace,
        release_data_root=release_data,
        slug=release.slug,
        artifact_id=PRESENTATION_ID,
    )

    (workspace / "course.yaml").write_text("not: [valid", encoding="utf-8")
    (workspace / "sources.yaml").write_text("version: 1\nsources: []\n", encoding="utf-8")
    (profile_dir(templates, profile.id) / "profile-v3.yaml").write_text(
        "not: [valid", encoding="utf-8"
    )
    changed_template = PPTXPresentation()
    changed_template.core_properties.title = "Changed current template"
    changed_template.save(str(profile_dir(templates, profile.id) / "template.pptx"))

    assert (
        regenerate_release_artifact(
            workspace=workspace,
            release_data_root=release_data,
            slug=release.slug,
            artifact_id=PRESENTATION_ID,
        )
        == original
    )

    # Regeneration is a recovery operation, so it must not depend on the
    # generated Artifact still being present in application data.
    (
        release_data / release.course_id / release.slug / "artifacts" / f"{PRESENTATION_ID}.pptx"
    ).unlink()
    assert (
        regenerate_release_artifact(
            workspace=workspace,
            release_data_root=release_data,
            slug=release.slug,
            artifact_id=PRESENTATION_ID,
        )
        == original
    )


def test_forged_or_noncanonical_release_manifests_are_rejected(tmp_path: Path) -> None:
    workspace = tmp_path / "course"
    _clean_course(workspace)
    release = publish_release(
        workspace=workspace,
        release_data_root=tmp_path / "releases",
        templates_data_root=tmp_path / "templates",
        request=PublishReleaseRequest(
            slug="fall-2026",
            name="Fall 2026",
            selection=ReleaseSelection(lecture_ids=[LECTURE_ID], artifact_ids=[PRESENTATION_ID]),
        ),
    )
    payload = release.model_dump(mode="json")
    payload["artifacts"][0]["storage_path"] = "artifacts/presentation-bbbbbbbbbbbb.pptx"
    forged = (
        json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
        + b"\n"
    )
    _replace_release_tag(workspace, release.slug, release.commit_oid, forged)
    with pytest.raises(ReleaseError, match="invalid Artifact path"):
        read_release(workspace, release.slug)

    payload = release.model_dump(mode="json")
    noncanonical = (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode()
    _replace_release_tag(workspace, release.slug, release.commit_oid, noncanonical)
    with pytest.raises(ReleaseError, match="not exact canonical JSON"):
        read_release(workspace, release.slug)
