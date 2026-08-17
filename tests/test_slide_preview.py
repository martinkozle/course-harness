from __future__ import annotations

import os
import re
import shutil
import subprocess
from pathlib import Path

import httpx2
import pytest
from pptx import Presentation as PPTXPresentation

from course_harness.app import create_app
from course_harness.course_plan import (
    CoursePlanInput,
    LectureInput,
    create_course_plan,
    create_course_plan_file,
    initialize_workspace_history,
)
from course_harness.presentation import BulletsSlide, Presentation
from course_harness.slide_preview import (
    PreviewContext,
    build_preview,
    normalized_layout_slots,
    render_presentation_preview,
    renderer_capability,
)
from course_harness.template_inspect import inspect_template
from course_harness.template_profiles import (
    TemplateLayoutMapping,
    TemplateProfile,
    profile_dir,
    write_profile_version,
)


def _profile(template_path: Path) -> TemplateProfile:
    inspection = inspect_template(template_path)
    return TemplateProfile(
        id="tpl-000000000001",
        name="Preview template",
        version=2,
        template_filename=template_path.name,
        slide_width=inspection["slide_width"],
        slide_height=inspection["slide_height"],
        slide_count=inspection["slide_count"],
        layouts=[
            TemplateLayoutMapping(
                semantic_layout="bullets",
                template_layout_index=1,
                confidence=1,
                rationale="Test mapping",
                slot_mappings={"title": 0, "body": 1},
            )
        ],
    )


def test_layout_slots_are_normalized_and_keep_extracted_typography(tmp_path: Path) -> None:
    template_path = tmp_path / "template.pptx"
    PPTXPresentation().save(str(template_path))
    profile = _profile(template_path)

    inspection = inspect_template(template_path)
    slots = normalized_layout_slots(profile, inspection, "bullets")

    assert set(slots) == {"title", "body"}
    assert slots["title"].left == 0.05
    assert 0 < slots["body"].width <= 1
    assert slots["title"].font_size > 0
    assert slots["title"].font_family
    assert slots["title"].font_family == inspection["theme"]["fonts"]["major"]


def test_preview_keeps_semantic_fallback_when_renderer_is_unavailable(
    tmp_path: Path, monkeypatch
) -> None:
    template_path = tmp_path / "template.pptx"
    PPTXPresentation().save(str(template_path))
    profile = _profile(template_path)
    presentation = Presentation(
        id="presentation-abc123def456",
        lecture_id="lecture-abc123def456",
        slides=[
            BulletsSlide(
                id="slide-abc123def456",
                title="Visible without LibreOffice",
                bullets=["One", "Two"],
            )
        ],
    )
    monkeypatch.setattr("course_harness.slide_preview.shutil.which", lambda _name: None)

    preview = build_preview(
        presentation, PreviewContext(profile, template_path, tmp_path / "cache")
    )

    assert preview.renderer.available is False
    assert preview.slides[0].thumbnail_url is None
    assert preview.slides[0].slots["body"].top > 0
    assert preview.slides[0].background_url is None


def test_renderer_capability_reports_detected_binary(monkeypatch) -> None:
    monkeypatch.setattr(
        "course_harness.slide_preview.shutil.which",
        lambda _name: "/nix/store/libreoffice/bin/libreoffice",
    )

    capability = renderer_capability()

    assert capability.available is True
    assert capability.name == "LibreOffice"


def test_render_populates_cached_background_and_authoritative_thumbnail(
    tmp_path: Path, monkeypatch
) -> None:
    template_path = tmp_path / "template.pptx"
    PPTXPresentation().save(str(template_path))
    profile = _profile(template_path)
    presentation = Presentation(
        id="presentation-abc123def456",
        lecture_id="lecture-abc123def456",
        slides=[
            BulletsSlide(
                id="slide-abc123def456",
                title="Rendered title",
                bullets=["Rendered body"],
            ),
            BulletsSlide(
                id="slide-abc123def457",
                title="Archived title",
                bullets=["Archived body"],
                archived=True,
            ),
        ],
    )
    monkeypatch.setattr(
        "course_harness.slide_preview.shutil.which", lambda _name: "/bin/libreoffice"
    )

    def fake_render(_pptx_bytes: bytes, target: Path) -> None:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b"png")

    monkeypatch.setattr("course_harness.slide_preview._render_pptx_png", fake_render)

    rendered = render_presentation_preview(
        presentation, PreviewContext(profile, template_path, tmp_path / "cache")
    )

    assert rendered.slides[0].background_url is not None
    assert rendered.slides[0].thumbnail_url is not None
    assert rendered.render_key in rendered.slides[0].thumbnail_url
    assert rendered.slides[1].thumbnail_url is not None


def test_libreoffice_adapter_produces_png_without_pixel_identity_requirement(
    tmp_path: Path,
) -> None:
    if not renderer_capability().available:
        pytest.skip("LibreOffice is not installed")
    template_path = tmp_path / "template.pptx"
    PPTXPresentation().save(str(template_path))
    profile = _profile(template_path)
    presentation = Presentation(
        id="presentation-abc123def456",
        lecture_id="lecture-abc123def456",
        slides=[
            BulletsSlide(
                id="slide-abc123def456",
                title="Adapter smoke",
                bullets=["Renderer output stays editable and readable"],
            )
        ],
    )

    rendered = render_presentation_preview(
        presentation, PreviewContext(profile, template_path, tmp_path / "cache")
    )

    thumbnail = (
        tmp_path
        / "cache"
        / profile.id
        / "previews"
        / rendered.render_key
        / "slide-abc123def456.png"
    )
    png = thumbnail.read_bytes()
    assert png.startswith(b"\x89PNG\r\n\x1a\n")
    width = int.from_bytes(png[16:20], "big")
    height = int.from_bytes(png[20:24], "big")
    assert width >= 500
    assert width / height == pytest.approx(profile.slide_width / profile.slide_height, rel=0.03)

    golden = Path(__file__).parent / "fixtures" / "previews" / "libreoffice-bullets.png"
    if os.environ.get("COURSE_HARNESS_UPDATE_PREVIEW_FIXTURE") == "1":
        golden.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(thumbnail, golden)
    comparator = shutil.which("compare")
    if comparator is None:
        pytest.skip("ImageMagick compare is not installed")
    assert golden.is_file(), "Run with COURSE_HARNESS_UPDATE_PREVIEW_FIXTURE=1"
    comparison = subprocess.run(
        [comparator, "-metric", "RMSE", str(golden), str(thumbnail), "null:"],
        capture_output=True,
        text=True,
        check=False,
    )
    assert comparison.returncode in {0, 1}
    normalized_match = re.search(r"\((\d+(?:\.\d+)?)\)", comparison.stderr)
    assert normalized_match is not None
    assert float(normalized_match.group(1)) <= 0.06


@pytest.mark.anyio
async def test_preview_api_exposes_semantic_fallback_and_renderer_state(
    tmp_path: Path, monkeypatch
) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    plan = create_course_plan(
        CoursePlanInput(
            title="Preview course",
            audience="Course Authors",
            lectures=[LectureInput(title="Lecture one")],
        )
    ).model_copy(
        update={
            "template_profile_id": "tpl-000000000001",
            "template_profile_version": 2,
        }
    )
    initialize_workspace_history(workspace)
    create_course_plan_file(workspace, plan)
    template_data = tmp_path / "templates"
    template_path = profile_dir(template_data, "tpl-000000000001") / "template.pptx"
    template_path.parent.mkdir(parents=True)
    PPTXPresentation().save(str(template_path))
    profile = _profile(template_path)
    write_profile_version(template_data, profile)
    from course_harness.presentation import write_presentation

    write_presentation(
        workspace,
        Presentation(
            id="presentation-abc123def456",
            lecture_id=plan.lectures[0].id,
            slides=[
                BulletsSlide(
                    id="slide-abc123def456",
                    title="API preview",
                    bullets=["Semantic content"],
                )
            ],
        ),
    )
    monkeypatch.setattr("course_harness.slide_preview.shutil.which", lambda _name: None)
    app = create_app(
        workspace,
        templates_data_path=template_data,
        templates_cache_path=tmp_path / "cache",
    )

    async with httpx2.AsyncClient(
        transport=httpx2.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.get(f"/api/presentations/{plan.lectures[0].id}/preview")
        render = await client.post(f"/api/presentations/{plan.lectures[0].id}/preview/render")

    assert response.status_code == 200
    assert response.json()["slides"][0]["slots"]["title"]["left"] == 0.05
    assert response.json()["renderer"]["available"] is False
    assert render.status_code == 422
    assert "Semantic previews remain available" in render.json()["detail"]
