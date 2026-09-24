from __future__ import annotations

from io import BytesIO

from pptx import Presentation as PPTXPresentation
from pptx.oxml.ns import qn

from course_harness.export import export_presentation
from course_harness.presentation import (
    BigStatementSlide,
    BulletsSlide,
    ClosingSlide,
    Presentation,
)
from course_harness.slide_fit import TextSlot, fit_estimate, fitted_font_size, min_font_size
from course_harness.template_slots import (
    find_body_placeholder,
    find_subtitle_or_body_placeholder,
    find_title_placeholder,
)

LONG_BULLET = (
    "Gradient descent iteratively updates every model parameter in the direction that "
    "reduces the training loss the most"
)


def _export_single(slide) -> object:
    presentation = Presentation(
        id="presentation-abc123def456", lecture_id="lecture-abc123def456", slides=[slide]
    )
    exported = PPTXPresentation(BytesIO(export_presentation(presentation)))
    return exported.slides[0]


def _run_sizes(placeholder) -> set[float | None]:
    return {
        run.font.size.pt if run.font.size is not None else None
        for paragraph in placeholder.text_frame.paragraphs
        for run in paragraph.runs
    }


def _has_norm_autofit(placeholder) -> bool:
    body_pr = placeholder.text_frame._txBody.find(qn("a:bodyPr"))
    return body_pr.find(qn("a:normAutofit")) is not None and body_pr.get("wrap") == "square"


def test_long_bullets_export_with_reduced_font_and_shrink_on_overflow() -> None:
    slide = _export_single(
        BulletsSlide(id="slide-abc123def456", title="Too much", bullets=[LONG_BULLET] * 12)
    )

    body = find_body_placeholder(slide)
    sizes = _run_sizes(body)

    assert len(sizes) == 1
    size = sizes.pop()
    assert size is not None
    assert 12 <= size < 32
    assert _has_norm_autofit(body)


def test_short_bullets_keep_the_template_font_size() -> None:
    slide = _export_single(
        BulletsSlide(id="slide-abc123def456", title="Fits", bullets=["One idea", "Another idea"])
    )

    body = find_body_placeholder(slide)

    assert _run_sizes(body) == {None}
    assert _has_norm_autofit(body)


def test_monospace_estimate_needs_more_room_than_proportional_text() -> None:
    slot = TextSlot(width_pt=648, height_pt=356, font_size_pt=32)
    code = [f"total_{index:02d} = accumulate(total_{index:02d}, step)" for index in range(12)]

    assert fit_estimate(slot, code, monospace=True) > fit_estimate(slot, code)
    size = fitted_font_size(slot, code, monospace=True)
    assert size is not None and min_font_size(32) <= size < 32


def test_closing_and_statement_text_are_fitted_on_export() -> None:
    long_text = " ".join([LONG_BULLET] * 6)
    statement = _export_single(BigStatementSlide(id="slide-abc123def456", statement=long_text))
    closing = _export_single(ClosingSlide(id="slide-abc123def456", title="Thanks", text=long_text))

    statement_ph = find_title_placeholder(statement)
    closing_ph = find_subtitle_or_body_placeholder(closing)
    assert None not in _run_sizes(statement_ph)
    assert _has_norm_autofit(statement_ph)
    assert None not in _run_sizes(closing_ph)


def test_fit_estimate_ratio_grows_with_text_and_shrinks_with_font() -> None:
    slot = TextSlot(width_pt=648, height_pt=356, font_size_pt=32)

    assert fit_estimate(slot, ["Short"]) < 1
    assert fit_estimate(slot, [LONG_BULLET] * 12) > 1
    assert fit_estimate(slot, [LONG_BULLET] * 12, font_size_pt=16) < fit_estimate(
        slot, [LONG_BULLET] * 12
    )


def test_fitted_font_size_never_goes_under_the_floor() -> None:
    slot = TextSlot(width_pt=300, height_pt=100, font_size_pt=24)

    assert fitted_font_size(slot, ["Fits"]) is None
    assert fitted_font_size(slot, [LONG_BULLET] * 40) == min_font_size(24) == 14
    assert min_font_size(18) == 12
    assert min_font_size(10) == 10
