from __future__ import annotations

from contextlib import suppress
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING

from pptx import Presentation as PPTXPresentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from course_harness.presentation import (
    BigStatementSlide,
    BulletsSlide,
    ClosingSlide,
    CodeSlide,
    ImageSlide,
    Presentation,
    QuoteSlide,
    SectionSlide,
    SlideCitation,
    TitleSlide,
    TwoColumnSlide,
    strip_list_marker,
)
from course_harness.resources import image_dimensions
from course_harness.slide_fit import fit_placeholder_text
from course_harness.template_slots import (
    CONTENT_LIKE_TYPES,
    find_body_placeholder,
    find_column_placeholder_groups,
    find_content_placeholders,
    find_placeholder_of_type,
    find_slot_placeholder,
    find_slot_placeholder_or,
    find_subtitle_or_body_placeholder,
    find_title_placeholder,
)

if TYPE_CHECKING:
    from course_harness.sources import SourceImage, SourceImageResolver
    from course_harness.template_profiles import TemplateProfile

DEFAULT_LAYOUT_MAPPING: dict[str, int] = {
    "title": 0,
    "section": 2,
    "bullets": 1,
    "two_column": 3,
    "big_statement": 5,
    "closing": 2,
    "code": 1,
    "image": 8,
    "quote": 1,
}

CITATION_FONT_SIZE = Pt(9)
CITATION_BOX_HEIGHT = Inches(0.45)
CITATION_BOX_TOP = Inches(7.0)
SLIDE_WIDTH = Inches(13.333)


class ExportError(Exception):
    pass


def export_presentation(
    presentation: Presentation,
    profile: TemplateProfile | None = None,
    template_path: Path | None = None,
    image_resolver: SourceImageResolver | None = None,
) -> bytes:
    if profile is not None and profile.id != "_builtin-default":
        layout_mapping = {m.semantic_layout: m.template_layout_index for m in profile.layouts}
        slot_mapping = {m.semantic_layout: m.slot_mappings for m in profile.layouts}
    else:
        layout_mapping = DEFAULT_LAYOUT_MAPPING
        slot_mapping = {}

    if template_path is not None and template_path.exists():
        prs = PPTXPresentation(str(template_path))
        ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        sldIdLst = prs.slides._sldIdLst
        while len(sldIdLst) > 0:
            rId = sldIdLst[0].get(ns)
            if rId is not None:
                prs.part.drop_rel(rId)
            sldIdLst.remove(sldIdLst[0])
    else:
        prs = PPTXPresentation()

    active_slides = [s for s in presentation.slides if not s.archived]

    for slide in active_slides:
        layout_index = layout_mapping.get(slide.layout)
        if layout_index is None:
            raise ExportError(f"Unsupported slide layout: {slide.layout}")

        slide_layouts = prs.slide_layouts
        if layout_index >= len(slide_layouts):
            raise ExportError(
                f"Slide layout index {layout_index} is out of range "
                f"(only {len(slide_layouts)} layouts available)"
            )

        pptx_slide = prs.slides.add_slide(slide_layouts[layout_index])
        if isinstance(slide, ImageSlide):
            image = (
                image_resolver(slide.image_source_id)
                if image_resolver is not None and slide.image_source_id is not None
                else None
            )
            _populate_image_slide(slide, pptx_slide, slot_mapping.get("image", {}), image)
        else:
            _populate_slide(slide, pptx_slide, slot_mapping.get(slide.layout, {}))
        _add_citations(slide, pptx_slide)
        _add_speaker_notes(slide, pptx_slide)

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _populate_slide(slide, pptx_slide, slot_mappings: dict[str, int]) -> None:
    handlers = {
        "title": _populate_title_slide,
        "section": _populate_section_slide,
        "bullets": _populate_bullets_slide,
        "two_column": _populate_two_column_slide,
        "big_statement": _populate_big_statement_slide,
        "closing": _populate_closing_slide,
        "code": _populate_code_slide,
        "quote": _populate_quote_slide,
    }
    handler = handlers.get(slide.layout)
    if handler:
        handler(slide, pptx_slide, slot_mappings)


def _populate_title_slide(slide: TitleSlide, pptx_slide, slot_mappings: dict[str, int]) -> None:
    title_ph = find_slot_placeholder_or(pptx_slide, slot_mappings, "title", find_title_placeholder)
    subtitle_ph = find_slot_placeholder_or(
        pptx_slide, slot_mappings, "subtitle", find_subtitle_or_body_placeholder
    )
    if title_ph and slide.title:
        title_ph.text_frame.text = slide.title
        fit_placeholder_text(title_ph)
    if subtitle_ph and slide.subtitle:
        subtitle_ph.text_frame.text = slide.subtitle
        fit_placeholder_text(subtitle_ph)


def _populate_section_slide(slide: SectionSlide, pptx_slide, slot_mappings: dict[str, int]) -> None:
    title_ph = find_slot_placeholder_or(pptx_slide, slot_mappings, "title", find_title_placeholder)
    if title_ph and slide.title:
        title_ph.text_frame.text = slide.title


def _populate_bullets_slide(slide: BulletsSlide, pptx_slide, slot_mappings: dict[str, int]) -> None:
    title_ph = find_slot_placeholder_or(pptx_slide, slot_mappings, "title", find_title_placeholder)
    body_ph = find_slot_placeholder_or(pptx_slide, slot_mappings, "body", find_body_placeholder)
    if title_ph and slide.title:
        title_ph.text_frame.text = slide.title
    if body_ph and slide.bullets:
        tf = body_ph.text_frame
        tf.clear()
        for i, bullet in enumerate(slide.bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = strip_list_marker(bullet)
            p.level = 0
        fit_placeholder_text(body_ph, bulleted=True)


def _populate_two_column_slide(
    slide: TwoColumnSlide, pptx_slide, slot_mappings: dict[str, int]
) -> None:
    title_ph = find_slot_placeholder_or(pptx_slide, slot_mappings, "title", find_title_placeholder)
    if title_ph and slide.title:
        title_ph.text_frame.text = slide.title
    left_ph = find_slot_placeholder(pptx_slide, slot_mappings, "left")
    right_ph = find_slot_placeholder(pptx_slide, slot_mappings, "right")
    left_items = _list_items(slide.left_content)
    right_items = _list_items(slide.right_content)
    if left_ph is not None or right_ph is not None:
        if left_ph is not None:
            _write_column(left_ph, left_items)
        if right_ph is not None:
            _write_column(right_ph, right_items)
        return
    groups = find_column_placeholder_groups(pptx_slide)

    for i, group in enumerate(groups):
        items = left_items if i == 0 else right_items
        if len(group) == 1:
            _write_column(group[0], items)
            continue
        # A heading placeholder above a body: the first item is the column heading.
        heading, body = (items[0], items[1:]) if items else ("", [])
        group[0].text_frame.text = heading
        if heading:
            fit_placeholder_text(group[0])
        for ph in group[1:]:
            _write_column(ph, body)
            body = []


def _list_items(items: list[str]) -> list[str]:
    return [cleaned for item in items if (cleaned := strip_list_marker(item))]


def _write_column(placeholder, lines: list[str]) -> None:
    tf = placeholder.text_frame
    tf.clear()
    for k, line in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.text = line
    if lines:
        fit_placeholder_text(placeholder, bulleted=True)


def _populate_big_statement_slide(
    slide: BigStatementSlide, pptx_slide, slot_mappings: dict[str, int]
) -> None:
    title_ph = find_slot_placeholder_or(
        pptx_slide, slot_mappings, "statement", find_title_placeholder
    )
    if title_ph and slide.statement:
        title_ph.text_frame.text = slide.statement
        for paragraph in title_ph.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
        fit_placeholder_text(title_ph)


def _populate_closing_slide(slide: ClosingSlide, pptx_slide, slot_mappings: dict[str, int]) -> None:
    title_ph = find_slot_placeholder_or(pptx_slide, slot_mappings, "title", find_title_placeholder)
    subtitle_ph = find_slot_placeholder_or(
        pptx_slide, slot_mappings, "body", find_subtitle_or_body_placeholder
    )
    if title_ph and slide.title:
        title_ph.text_frame.text = slide.title
    if subtitle_ph and slide.text:
        subtitle_ph.text_frame.text = slide.text
        fit_placeholder_text(subtitle_ph)


def _populate_code_slide(slide: CodeSlide, pptx_slide, slot_mappings: dict[str, int]) -> None:
    title_ph = find_slot_placeholder_or(pptx_slide, slot_mappings, "title", find_title_placeholder)
    body_ph = find_slot_placeholder_or(pptx_slide, slot_mappings, "body", find_body_placeholder)
    if title_ph and slide.title:
        title_ph.text_frame.text = slide.title
    if body_ph and slide.code:
        tf = body_ph.text_frame
        tf.clear()
        header = slide.language or "code"
        p = tf.paragraphs[0]
        p.text = f"[{header}]"
        for line in slide.code.split("\n"):
            p = tf.add_paragraph()
            p.text = line
            for run in p.runs:
                run.font.name = "Courier New"


def _populate_image_slide(
    slide: ImageSlide,
    pptx_slide,
    slot_mappings: dict[str, int],
    image: SourceImage | None = None,
) -> None:
    title_ph = find_slot_placeholder_or(pptx_slide, slot_mappings, "title", find_title_placeholder)
    if title_ph and slide.title:
        title_ph.text_frame.text = slide.title
    mapped_image_ph = find_slot_placeholder(pptx_slide, slot_mappings, "image")
    if image is None:
        phs = (
            [mapped_image_ph]
            if mapped_image_ph is not None
            else find_content_placeholders(pptx_slide)
        )
        if phs and slide.caption:
            phs[0].text_frame.text = slide.caption
        elif phs and slide.image_url:
            phs[0].text_frame.text = slide.image_url
        return

    content_phs = find_content_placeholders(pptx_slide)
    image_ph = (
        mapped_image_ph
        or find_placeholder_of_type(pptx_slide, {PP_PLACEHOLDER.PICTURE})
        or next(iter(content_phs), None)
    )
    caption_ph = next((ph for ph in content_phs if ph is not image_ph), None)
    if caption_ph is not None and slide.caption:
        caption_ph.text_frame.text = slide.caption
    _place_image(pptx_slide, image_ph, image)


def _place_image(pptx_slide, placeholder, image: SourceImage) -> None:
    """Fit an image inside its placeholder's frame without cropping it."""
    content = image.path.read_bytes()
    pixel_width, pixel_height = image_dimensions(content)
    if placeholder is not None:
        left, top, width, height = (
            placeholder.left,
            placeholder.top,
            placeholder.width,
            placeholder.height,
        )
    else:
        left, top, width, height = Inches(1), Inches(1.5), Inches(8), Inches(5)
    scale = min(width / pixel_width, height / pixel_height)
    fitted_width, fitted_height = int(pixel_width * scale), int(pixel_height * scale)
    fitted_left = int(left + (width - fitted_width) / 2)
    fitted_top = int(top + (height - fitted_height) / 2)

    if placeholder is not None and hasattr(placeholder, "insert_picture"):
        picture = placeholder.insert_picture(BytesIO(content))
        picture.crop_left = picture.crop_right = picture.crop_top = picture.crop_bottom = 0
        picture.left, picture.top = fitted_left, fitted_top
        picture.width, picture.height = fitted_width, fitted_height
        return
    if placeholder is not None:
        placeholder.element.getparent().remove(placeholder.element)
    pptx_slide.shapes.add_picture(
        BytesIO(content), fitted_left, fitted_top, fitted_width, fitted_height
    )


def _populate_quote_slide(slide: QuoteSlide, pptx_slide, slot_mappings: dict[str, int]) -> None:
    title_ph = find_slot_placeholder_or(pptx_slide, slot_mappings, "title", find_title_placeholder)
    body_ph = find_slot_placeholder_or(pptx_slide, slot_mappings, "body", find_body_placeholder)
    if title_ph and slide.title:
        title_ph.text_frame.text = slide.title
    if body_ph:
        tf = body_ph.text_frame
        tf.clear()
        if slide.quote:
            p = tf.paragraphs[0]
            p.text = slide.quote
            p.alignment = PP_ALIGN.CENTER
        if slide.attribution:
            p = tf.add_paragraph()
            run = p.add_run()
            run.text = f"— {slide.attribution}"
            run.font.italic = True
            p.alignment = PP_ALIGN.CENTER
        fit_placeholder_text(body_ph)


def _add_citations(slide, pptx_slide) -> None:
    if not slide.citations:
        return
    left = Inches(0.5)
    top = CITATION_BOX_TOP
    width = SLIDE_WIDTH - Inches(1.0)
    height = CITATION_BOX_HEIGHT
    textbox = pptx_slide.shapes.add_textbox(left, top, width, height)
    tf = textbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Sources: "
    for i, citation in enumerate(slide.citations):
        if i > 0:
            p.add_run().text = ", "
        run = p.add_run()
        run.text = citation.label
        run.font.size = CITATION_FONT_SIZE
        if citation.url:
            run.hyperlink.address = citation.url
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.size = CITATION_FONT_SIZE


def _add_speaker_notes(slide, pptx_slide) -> None:
    note_lines: list[str] = []
    if slide.speaker_notes:
        note_lines.append(slide.speaker_notes)
    if slide.citations:
        citation_lines = _citation_note_lines(slide.citations)
        if note_lines and citation_lines:
            note_lines.append("")
        note_lines.extend(citation_lines)
    if not note_lines:
        return
    try:
        notes = pptx_slide.notes_slide
        notes.notes_text_frame.text = "\n".join(note_lines)
    except Exception:
        pass


def _citation_note_lines(citations: list[SlideCitation]) -> list[str]:
    lines: list[str] = []
    for c in citations:
        parts = [c.label]
        if c.url:
            parts.append(c.url)
        if c.line_start is not None:
            parts.append(
                f"lines {c.line_start + 1}"
                + (f"–{c.line_end + 1}" if c.line_end is not None else "")
            )
        lines.append(" | ".join(parts))
    return lines


PLACEHOLDER_REQUIREMENTS: dict[str, dict[str, str]] = {
    "title": {"title": "TITLE or CENTER_TITLE"},
    "section": {"title": "TITLE or CENTER_TITLE"},
    "bullets": {"title": "TITLE or CENTER_TITLE", "body": "BODY or OBJECT"},
    "two_column": {"title": "TITLE or CENTER_TITLE"},
    "big_statement": {},
    "closing": {"title": "TITLE or CENTER_TITLE"},
    "code": {"title": "TITLE or CENTER_TITLE", "body": "BODY or OBJECT"},
    "image": {},
    "quote": {"title": "TITLE or CENTER_TITLE", "body": "BODY or OBJECT"},
}

TITLE_LIKE_TYPES = {1, 3}  # TITLE, CENTER_TITLE


def validate_export_mapping(
    profile: TemplateProfile,
    template_path: Path | None = None,
) -> list[dict]:
    issues: list[dict] = []
    if profile.id == "_builtin-default":
        return issues

    if template_path is None or not template_path.exists():
        issues.append({"level": "blocking", "message": "Template file not found"})
        return issues

    prs = PPTXPresentation(str(template_path))
    slide_layouts = prs.slide_layouts

    for m in profile.layouts:
        semantic = m.semantic_layout
        idx = m.template_layout_index
        if idx >= len(slide_layouts):
            issues.append(
                {
                    "level": "blocking",
                    "message": (
                        f"Layout '{semantic}' maps to index {idx}, "
                        f"but template only has {len(slide_layouts)} layouts"
                    ),
                }
            )
            continue

        layout = slide_layouts[idx]
        ph_by_idx: dict[int, int] = {}
        for ph in layout.placeholders:
            with suppress(Exception):
                ph_by_idx[ph.placeholder_format.idx] = (
                    int(ph.placeholder_format.type) if ph.placeholder_format.type is not None else 0
                )

        layout_name = layout.name or f"Layout {idx}"
        reqs = PLACEHOLDER_REQUIREMENTS.get(semantic, {})

        required_indices: dict[str, int | None] = {}
        for slot, ph_desc in reqs.items():
            if slot in m.slot_mappings:
                required_indices[slot] = m.slot_mappings[slot]
                continue
            fallback = (
                find_title_placeholder(layout)
                if ph_desc == "TITLE or CENTER_TITLE"
                else find_body_placeholder(layout)
            )
            required_indices[slot] = (
                fallback.placeholder_format.idx if fallback is not None else None
            )

        resolved_slots = dict(m.slot_mappings)
        for slot, placeholder_index in required_indices.items():
            if placeholder_index is not None:
                resolved_slots.setdefault(slot, placeholder_index)
        slots_by_placeholder: dict[int, list[str]] = {}
        for slot, placeholder_index in resolved_slots.items():
            slots_by_placeholder.setdefault(placeholder_index, []).append(slot)
        for placeholder_index, slots in slots_by_placeholder.items():
            if len(slots) > 1:
                sorted_slots = sorted(slots)
                issues.append(
                    {
                        "level": "blocking",
                        "message": (
                            f"'{semantic}' → '{layout_name}' (index {idx}): "
                            f"slots '{sorted_slots[0]}' and '{sorted_slots[1]}' use the same "
                            f"placeholder idx={placeholder_index}"
                        ),
                    }
                )

        for slot, ph_desc in reqs.items():
            placeholder_index = required_indices[slot]
            actual = ph_by_idx.get(placeholder_index) if placeholder_index is not None else None
            if actual is None:
                location = (
                    f"missing placeholder idx={placeholder_index}"
                    if placeholder_index is not None
                    else "no compatible placeholder"
                )
                issues.append(
                    {
                        "level": "blocking",
                        "message": (
                            f"'{semantic}' → '{layout_name}' (index {idx}): "
                            f"slot '{slot}' has {location} (expected {ph_desc})"
                        ),
                    }
                )
            elif ph_desc == "TITLE or CENTER_TITLE" and actual not in TITLE_LIKE_TYPES:
                issues.append(
                    {
                        "level": "blocking",
                        "message": (
                            f"'{semantic}' → '{layout_name}' (index {idx}): "
                            f"placeholder idx={placeholder_index} is type {actual}, "
                            f"expected {ph_desc}"
                        ),
                    }
                )
            elif ph_desc == "BODY or OBJECT" and actual not in CONTENT_LIKE_TYPES:
                issues.append(
                    {
                        "level": "blocking",
                        "message": (
                            f"'{semantic}' → '{layout_name}' (index {idx}): "
                            f"slot '{slot}' maps to type {PP_PLACEHOLDER(actual).name}, "
                            f"expected {ph_desc}"
                        ),
                    }
                )

        for slot in ("left", "right"):
            placeholder_index = m.slot_mappings.get(slot)
            if placeholder_index is None:
                continue
            actual = ph_by_idx.get(placeholder_index)
            if actual is not None and actual not in CONTENT_LIKE_TYPES:
                issues.append(
                    {
                        "level": "blocking",
                        "message": (
                            f"'{semantic}' → '{layout_name}' (index {idx}): "
                            f"slot '{slot}' maps to type {PP_PLACEHOLDER(actual).name}, "
                            "expected BODY or OBJECT"
                        ),
                    }
                )

        if semantic == "two_column":
            configured_columns = {
                slot: m.slot_mappings[slot] for slot in ("left", "right") if slot in m.slot_mappings
            }
            column_groups = find_column_placeholder_groups(layout)
            if configured_columns and len(configured_columns) != 2:
                issues.append(
                    {
                        "level": "blocking",
                        "message": (
                            f"'{semantic}' → '{layout_name}' (index {idx}): "
                            "both left and right slots must be configured"
                        ),
                    }
                )
            elif configured_columns and any(
                index not in ph_by_idx for index in configured_columns.values()
            ):
                issues.append(
                    {
                        "level": "blocking",
                        "message": (
                            f"'{semantic}' → '{layout_name}' (index {idx}): "
                            "a configured column slot does not exist"
                        ),
                    }
                )
            elif not configured_columns and len(column_groups) < 2:
                issues.append(
                    {
                        "level": "blocking",
                        "message": (
                            f"'{semantic}' → '{layout_name}' (index {idx}): "
                            f"only {len(column_groups)} content column group(s) found, "
                            f"two_column needs at least 2"
                        ),
                    }
                )

        if semantic == "image":
            has_pic = any(
                ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE for ph in layout.placeholders
            )
            if not has_pic and 0 not in ph_by_idx:
                issues.append(
                    {
                        "level": "warning",
                        "message": (
                            f"'{semantic}' → '{layout_name}' (index {idx}): "
                            "no PICTURE or TITLE placeholder found"
                        ),
                    }
                )

    return issues
