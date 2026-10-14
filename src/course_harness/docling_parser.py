"""Docling conversion for the local Document Processors, and full Docling's model provisioning.

The default install has Docling Slim without models: PDFs go through its native pipeline,
which reads the text and images already in the file. Full Docling, installed with the
``course-harness[docling]`` extra, adds layout, table, and OCR models (ADR 0015).
"""

import importlib.util
import io
import threading
import unicodedata
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pydantic import BaseModel

PDF_MEDIA_TYPE = "application/pdf"
PPTX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
OFFICE_MEDIA_TYPES = {PPTX_MEDIA_TYPE: ".pptx", DOCX_MEDIA_TYPE: ".docx"}
HTML_MEDIA_TYPES = {"text/html": ".html", "application/xhtml+xml": ".xhtml"}
MAX_DOCUMENT_BYTES = 100 * 1024 * 1024

# The full PDF pipeline runs layout, table, and OCR models that need gigabytes of memory
# and every core, so those PDFs convert one at a time. Nothing else waits for them.
_MODEL_PDF_CONVERSION = threading.Lock()


def full_docling_installed() -> bool:
    """Whether the ``course-harness[docling]`` extra, with its models runtime, is installed."""
    return all(
        importlib.util.find_spec(module) is not None for module in ("torch", "docling_ibm_models")
    )


def models_ready(cache_dir: Path) -> bool:
    return (cache_dir / "docling-models.ready").is_file()


ModelProgress = Callable[[str, int], None]


class ModelDownloadStatus(BaseModel):
    ready: bool
    # Whether full Docling is installed at all; the models are useless without it.
    installed: bool = True
    downloading: bool = False
    stage: str | None = None
    completed_steps: int = 0
    total_steps: int = 3
    error: str | None = None


def download_models(cache_dir: Path, on_progress: ModelProgress | None = None) -> None:
    """Prefetch only the layout, table, and OCR models used for PDF conversion."""
    from docling.utils.model_downloader import download_models as docling_download_models

    models_dir = cache_dir / "docling-models"
    stages = (
        "Downloading layout models",
        "Downloading table model",
        "Downloading text recognition models",
    )
    for completed, stage in enumerate(stages):
        if on_progress is not None:
            on_progress(stage, completed)
        docling_download_models(
            output_dir=models_dir,
            with_layout=completed == 0,
            with_tableformer=completed == 1,
            with_rapidocr=completed == 2,
            with_code_formula=False,
            with_picture_classifier=False,
        )
    (cache_dir / "docling-models.ready").write_text("ready\n", encoding="utf-8")
    if on_progress is not None:
        on_progress("Ready", len(stages))


@dataclass(frozen=True)
class PageContent:
    """What one page or slide of a converted document holds, for the quality check."""

    number: int
    characters: int
    garbled_characters: int
    # The share of the page covered by images, from 0 to 1.
    image_coverage: float


@dataclass(frozen=True)
class DoclingConversion:
    markdown: str
    structured: str
    pages: list[PageContent]


def convert_document(
    media_type: str, content: bytes, cache_dir: Path, *, use_models: bool = False
) -> DoclingConversion:
    """Convert a document to page-sectioned Markdown and the structured Docling JSON.

    ``use_models`` selects full Docling's model pipeline for a PDF; otherwise PDFs use
    the native pipeline, which runs no model at all.
    """
    from docling.datamodel.base_models import DocumentStream, InputFormat
    from docling.document_converter import DocumentConverter

    if media_type == PDF_MEDIA_TYPE:
        extension = ".pdf"
        converter = _model_pdf_converter(cache_dir) if use_models else _native_pdf_converter()
    elif media_type in HTML_MEDIA_TYPES:
        extension = HTML_MEDIA_TYPES[media_type]
        converter = DocumentConverter(allowed_formats=[InputFormat.HTML])
    else:
        extension = OFFICE_MEDIA_TYPES[media_type]
        converter = DocumentConverter(
            allowed_formats=[InputFormat.PPTX if extension == ".pptx" else InputFormat.DOCX]
        )

    stream = DocumentStream(name=f"resource{extension}", stream=io.BytesIO(content))
    if media_type == PDF_MEDIA_TYPE and use_models:
        with _MODEL_PDF_CONVERSION:
            result = converter.convert(stream, max_file_size=MAX_DOCUMENT_BYTES)
    else:
        result = converter.convert(stream, max_file_size=MAX_DOCUMENT_BYTES)
    document = result.document
    if document.pages:
        label = "Slide" if extension == ".pptx" else "Page"
        notes = _powerpoint_notes(content) if extension == ".pptx" else {}
        sections = [
            page_section(
                label,
                page_no,
                document.export_to_markdown(page_no=page_no),
                notes.get(page_no),
            )
            for page_no in sorted(document.pages)
        ]
        markdown = "\n\n".join(sections)
    else:
        markdown = document.export_to_markdown().strip()
    return DoclingConversion(
        markdown=markdown + ("\n" if markdown else ""),
        structured=document.model_dump_json(),
        pages=_page_contents(document),
    )


def page_section(label: str, number: int, body: str, notes: str | None = None) -> str:
    """One page or slide of a Derived Representation, in the shape every processor emits."""
    section = f"# {label} {number}\n\n{body.strip()}"
    if notes:
        section += f"\n\nNotes: {notes}"
    return section


def count_garbled(text: str) -> int:
    """Characters that only appear when a PDF's text could not be decoded."""
    garbled = text.count("(cid:") * 5
    for character in text:
        if character == "\ufffd" or (
            character not in "\n\t" and unicodedata.category(character) in {"Cc", "Co"}
        ):
            garbled += 1
    return garbled


def _native_pdf_converter() -> Any:
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import NativePdfPipelineOptions
    from docling.document_converter import DocumentConverter, NativePdfFormatOption

    options = NativePdfPipelineOptions(generate_page_images=False, generate_picture_images=False)
    return DocumentConverter(
        allowed_formats=[InputFormat.PDF],
        format_options={InputFormat.PDF: NativePdfFormatOption(pipeline_options=options)},
    )


def _model_pdf_converter(cache_dir: Path) -> Any:
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    if not full_docling_installed():
        raise ValueError("Full Docling is not installed.")
    if not models_ready(cache_dir):
        raise ValueError("PDF processing models have not been downloaded.")
    options = PdfPipelineOptions(artifacts_path=cache_dir / "docling-models")
    return DocumentConverter(
        allowed_formats=[InputFormat.PDF],
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)},
    )


def _page_contents(document: Any) -> list[PageContent]:
    characters: dict[int, int] = dict.fromkeys(document.pages, 0)
    garbled: dict[int, int] = dict.fromkeys(document.pages, 0)
    image_area: dict[int, float] = dict.fromkeys(document.pages, 0.0)
    for item in [*document.texts, *document.tables]:
        text = getattr(item, "text", None)
        if text is None:
            text = item.export_to_markdown(doc=document)
        for provenance in item.prov:
            if provenance.page_no in characters:
                characters[provenance.page_no] += len(text.strip())
                garbled[provenance.page_no] += count_garbled(text)
    for picture in document.pictures:
        for provenance in picture.prov:
            if provenance.page_no in image_area:
                image_area[provenance.page_no] += provenance.bbox.area()
    pages: list[PageContent] = []
    for number, page in sorted(document.pages.items()):
        page_area = page.size.width * page.size.height if page.size else 0
        coverage = min(1.0, image_area[number] / page_area) if page_area > 0 else 0.0
        pages.append(
            PageContent(
                number=number,
                characters=characters[number],
                garbled_characters=garbled[number],
                image_coverage=coverage,
            )
        )
    return pages


def _powerpoint_notes(content: bytes) -> dict[int, str]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(content))
    notes: dict[int, str] = {}
    for page_no, slide in enumerate(presentation.slides, start=1):
        if slide.has_notes_slide:
            note_frame = slide.notes_slide.notes_text_frame
            if note_frame is not None and note_frame.text.strip():
                notes[page_no] = note_frame.text.strip()
    return notes
