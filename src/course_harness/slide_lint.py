"""Check authored Slides against the active template before anyone renders them.

The linter exports the Slides with the real exporter, so it sees the same
placeholders, citation strip and text as a PowerPoint export, and then measures
each filled placeholder with the text-fit heuristic at the smallest font size the
exporter may use.  It needs no LibreOffice and gives the same answer every time.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING

from pptx import Presentation as PPTXPresentation

from course_harness.export import export_presentation
from course_harness.presentation import Presentation, Slide, strip_list_marker
from course_harness.slide_fit import (
    BULLET_INDENT_PT,
    fit_estimate,
    min_font_size,
    placeholder_font_size,
    placeholder_slot,
)

if TYPE_CHECKING:
    from course_harness.template_profiles import TemplateProfile

MAX_BULLETS = 6
MAX_COLUMN_ITEMS = 6  # including the heading item
MAX_ITEM_WORDS = 20
MAX_TITLE_WORDS = 6
MAX_STATEMENT_WORDS = 25
MAX_QUOTE_WORDS = 40
MAX_CODE_LINES = 15

# Content limits the Course Agent is told about, per layout.
CONTENT_LIMITS_GUIDE = (
    f"Slide titles should be concise (1-{MAX_TITLE_WORDS} words). "
    f"bullets: at most {MAX_BULLETS} bullets of up to {MAX_ITEM_WORDS} words each; "
    f"two_column: at most {MAX_COLUMN_ITEMS} items per column, each up to {MAX_ITEM_WORDS} "
    f"words; big_statement: one statement of up to {MAX_STATEMENT_WORDS} words; "
    f"quote: up to {MAX_QUOTE_WORDS} words; code: up to {MAX_CODE_LINES} short lines. "
    "Split a Slide instead of cramming it."
)

# The fields a layout cannot do without.  image Slides need an image Source.
_REQUIRED_FIELDS: dict[str, tuple[str, ...]] = {
    "title": ("title",),
    "section": ("title",),
    "bullets": ("title", "bullets"),
    "two_column": ("title", "left_content", "right_content"),
    "big_statement": ("statement",),
    "closing": ("title",),
    "code": ("code",),
    "image": ("image_source_id",),
    "quote": ("quote",),
}
_LIST_FIELDS = ("bullets", "left_content", "right_content")
_TEXT_FIELDS = (
    "title",
    "subtitle",
    *_LIST_FIELDS,
    "statement",
    "text",
    "code",
    "caption",
    "quote",
    "attribution",
)


@dataclass(frozen=True)
class SlideFinding:
    slide_id: str
    code: str
    field: str
    message: str

    def __str__(self) -> str:
        return f"{self.slide_id} {self.field} {self.message}"


def lint_slides(
    slides: list[Slide],
    profile: TemplateProfile | None = None,
    template_path: Path | None = None,
) -> list[SlideFinding]:
    """Findings for the active Slides given, in Slide order."""
    active = [slide for slide in slides if not slide.archived]
    if not active:
        return []
    deck = Presentation(id="presentation-000000000000", lecture_id="lecture-lint", slides=active)
    exported = PPTXPresentation(BytesIO(export_presentation(deck, profile, template_path)))
    findings: list[SlideFinding] = []
    for slide, pptx_slide in zip(active, exported.slides, strict=True):
        findings.extend(_content_findings(slide))
        findings.extend(_overflow_findings(slide, pptx_slide))
    return findings


def format_findings(findings: list[SlideFinding]) -> str:
    if not findings:
        return "Layout check: every Slide fits its template."
    lines = "\n".join(f"- {finding}" for finding in findings)
    return f"Layout check found {len(findings)} problem(s):\n{lines}"


def _content_findings(slide: Slide) -> list[SlideFinding]:
    findings: list[SlideFinding] = []

    def add(code: str, field: str, message: str) -> None:
        findings.append(SlideFinding(slide.id, code, field, message))

    for field in _REQUIRED_FIELDS.get(slide.layout, ()):
        value = getattr(slide, field, None)
        if field == "image_source_id" and getattr(slide, "image_url", None):
            continue
        if not (value.strip() if isinstance(value, str) else value):
            add("empty-slot", field, f"is empty — the {slide.layout} layout needs it")

    title = getattr(slide, "title", None) or ""
    if _words(title) > MAX_TITLE_WORDS:
        add("long-title", "title", f"has {_words(title)} words — keep it to {MAX_TITLE_WORDS}")

    for field in _LIST_FIELDS:
        items: list[str] = getattr(slide, field, None) or []
        limit = MAX_BULLETS if field == "bullets" else MAX_COLUMN_ITEMS
        if len(items) > limit:
            add(
                "too-many-items",
                field,
                f"has {len(items)} items — keep at most {limit} or split the Slide",
            )
        long_items = [
            number for number, item in enumerate(items, start=1) if _words(item) > MAX_ITEM_WORDS
        ]
        if long_items:
            numbers = ", ".join(str(number) for number in long_items)
            add(
                "long-item",
                field,
                f"item{'s' if len(long_items) > 1 else ''} {numbers} "
                f"{'are' if len(long_items) > 1 else 'is'} over {MAX_ITEM_WORDS} words — "
                "shorten them and move detail to speaker_notes",
            )

    for field, limit in (("statement", MAX_STATEMENT_WORDS), ("quote", MAX_QUOTE_WORDS)):
        text = getattr(slide, field, None) or ""
        if _words(text) > limit:
            add("long-text", field, f"has {_words(text)} words — keep it to {limit}")

    code = getattr(slide, "code", None) or ""
    code_lines = len(code.splitlines())
    if code_lines > MAX_CODE_LINES:
        add(
            "long-code",
            "code",
            f"has {code_lines} lines — show at most {MAX_CODE_LINES} or split the Slide",
        )
    return findings


def _overflow_findings(slide: Slide, pptx_slide) -> list[SlideFinding]:
    findings: list[SlideFinding] = []
    for placeholder in pptx_slide.placeholders:
        if not placeholder.has_text_frame:
            continue
        paragraphs = [paragraph.text for paragraph in placeholder.text_frame.paragraphs]
        if not any(text.strip() for text in paragraphs):
            continue
        field = _field_for(slide, paragraphs)
        monospace = field == "code"
        smallest = min_font_size(placeholder_font_size(placeholder))
        ratio = fit_estimate(
            placeholder_slot(placeholder),
            paragraphs,
            font_size_pt=smallest,
            monospace=monospace,
            indent_pt=BULLET_INDENT_PT if field in _LIST_FIELDS else 0.0,
        )
        if ratio <= 1.0:
            continue
        if monospace:
            target = max(1, math.floor(len(paragraphs) / ratio))
            advice = f"cut to ~{target} lines or split"
        else:
            words = sum(_words(text) for text in paragraphs)
            target = max(1, math.floor(words / ratio))
            advice = f"cut to ~{target} words or split"
        findings.append(
            SlideFinding(
                slide.id,
                "overflow",
                field,
                f"~{round(ratio * 100)}% of box even at min size — {advice}",
            )
        )
    return findings


def _field_for(slide: Slide, paragraphs: list[str]) -> str:
    """Name the Slide field whose text fills a placeholder."""
    first = next(text.strip() for text in paragraphs if text.strip())
    for field in _TEXT_FIELDS:
        value = getattr(slide, field, None)
        if not value:
            continue
        lines = value if isinstance(value, list) else value.splitlines()
        if first in {strip_list_marker(line) for line in lines} or first in {
            line.strip() for line in lines
        }:
            return field
        if field == "attribution" and first.endswith(value):
            return field
    return "text"


def _words(text: str) -> int:
    return len(text.split())
