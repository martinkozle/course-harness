from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import tempfile
import threading
from collections.abc import Callable
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from pptx import Presentation as PPTXPresentation
from pydantic import BaseModel, ConfigDict

from course_harness.export import export_presentation
from course_harness.presentation import Presentation, Slide
from course_harness.template_inspect import infer_slot_mappings, inspect_template
from course_harness.template_profiles import TemplateProfile


@dataclass(frozen=True)
class PreviewContext:
    profile: TemplateProfile
    template_path: Path | None
    cache_dir: Path


class RendererCapability(BaseModel):
    model_config = ConfigDict(extra="forbid")

    available: bool
    name: str
    detail: str


class PreviewSlot(BaseModel):
    model_config = ConfigDict(extra="forbid")

    left: float
    top: float
    width: float
    height: float
    font_family: str
    font_size: float
    bold: bool
    alignment: str | None


class SlidePreview(BaseModel):
    model_config = ConfigDict(extra="forbid")

    slide_id: str
    layout: str
    render_key: str
    slots: dict[str, PreviewSlot]
    background_url: str | None
    thumbnail_url: str | None


class PresentationPreview(BaseModel):
    model_config = ConfigDict(extra="forbid")

    profile_id: str
    profile_version: int
    slide_width: int
    slide_height: int
    renderer: RendererCapability
    render_key: str
    slides: list[SlidePreview]


def _libreoffice_executable() -> str | None:
    return shutil.which("libreoffice") or shutil.which("soffice")


def renderer_capability() -> RendererCapability:
    executable = _libreoffice_executable()
    if executable:
        return RendererCapability(
            available=True,
            name="LibreOffice",
            detail="High-fidelity thumbnails are available.",
        )
    return RendererCapability(
        available=False,
        name="LibreOffice",
        detail="LibreOffice was not found. Semantic previews remain available.",
    )


def normalized_layout_slots(
    profile: TemplateProfile,
    inspection: dict,
    semantic_layout: str,
) -> dict[str, PreviewSlot]:
    mapping = next(
        (item for item in profile.layouts if item.semantic_layout == semantic_layout), None
    )
    if mapping is None or mapping.template_layout_index >= len(inspection["layouts"]):
        return {}
    layout = inspection["layouts"][mapping.template_layout_index]
    placeholders = {item["idx"]: item for item in layout["placeholders"]}
    result: dict[str, PreviewSlot] = {}
    slot_mappings = mapping.slot_mappings or infer_slot_mappings(layout, semantic_layout)
    for slot, placeholder_index in slot_mappings.items():
        placeholder = placeholders.get(placeholder_index)
        if placeholder is None:
            continue
        result[slot] = PreviewSlot(
            left=round(placeholder["left"] / profile.slide_width, 6),
            top=round(placeholder["top"] / profile.slide_height, 6),
            width=round(placeholder["width"] / profile.slide_width, 6),
            height=round(placeholder["height"] / profile.slide_height, 6),
            font_family=placeholder["font_family"],
            font_size=placeholder["font_size"],
            bold=placeholder["bold"],
            alignment=placeholder["alignment"],
        )
    return result


_PREVIEW_CACHE_VERSION = 2
_RENDER_LOCK = threading.Lock()


def slide_render_key(slide: Slide, profile: TemplateProfile) -> str:
    mapping = next((item for item in profile.layouts if item.semantic_layout == slide.layout), None)
    payload = {
        "cache_version": _PREVIEW_CACHE_VERSION,
        "profile_id": profile.id,
        "profile_version": profile.version,
        "mapping": (
            {
                "semantic_layout": mapping.semantic_layout,
                "template_layout_index": mapping.template_layout_index,
                "slot_mappings": mapping.slot_mappings,
            }
            if mapping is not None
            else None
        ),
        "slide": slide.model_dump(
            mode="json",
            exclude={"id", "archived", "speaker_notes", "purpose"},
        ),
    }
    encoded = json.dumps(payload, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()[:20]


def presentation_render_key(presentation: Presentation, profile: TemplateProfile) -> str:
    payload = [slide_render_key(slide, profile) for slide in presentation.slides]
    return hashlib.sha256("".join(payload).encode("utf-8")).hexdigest()[:20]


def _render_cached_png(target: Path, make_pptx: Callable[[], bytes]) -> None:
    if target.is_file():
        return
    with _RENDER_LOCK:
        if target.is_file():
            return
        _render_pptx_png(make_pptx(), target)


def build_preview(
    presentation: Presentation,
    context: PreviewContext,
) -> PresentationPreview:
    profile = context.profile
    template_path = context.template_path
    cache_dir = context.cache_dir
    inspection = (
        inspect_template(template_path) if template_path is not None else _builtin_inspection()
    )
    render_key = presentation_render_key(presentation, profile)
    background_dir = cache_dir / profile.id / "backgrounds" / f"v{profile.version}"
    slides = []
    for slide in presentation.slides:
        slide_key = slide_render_key(slide, profile)
        thumbnail = cache_dir / profile.id / "previews" / slide_key / "thumbnail.png"
        background = background_dir / f"{slide.layout}.png"
        slides.append(
            SlidePreview(
                slide_id=slide.id,
                layout=slide.layout,
                render_key=slide_key,
                slots=normalized_layout_slots(profile, inspection, slide.layout),
                background_url=(
                    f"/api/preview-assets/{profile.id}/backgrounds/v{profile.version}/{slide.layout}.png"
                    if background.is_file()
                    else None
                ),
                thumbnail_url=(
                    f"/api/preview-assets/{profile.id}/previews/{slide_key}/thumbnail.png"
                    if thumbnail.is_file()
                    else None
                ),
            )
        )
    return PresentationPreview(
        profile_id=profile.id,
        profile_version=profile.version,
        slide_width=profile.slide_width,
        slide_height=profile.slide_height,
        renderer=renderer_capability(),
        render_key=render_key,
        slides=slides,
    )


def render_presentation_preview(
    presentation: Presentation,
    context: PreviewContext,
) -> PresentationPreview:
    profile = context.profile
    template_path = context.template_path
    cache_dir = context.cache_dir
    capability = renderer_capability()
    if not capability.available:
        raise RuntimeError(capability.detail)
    render_layout_backgrounds(context)
    for slide in presentation.slides:
        slide_key = slide_render_key(slide, profile)
        target = cache_dir / profile.id / "previews" / slide_key / "thumbnail.png"

        def make_pptx(slide=slide) -> bytes:
            renderable_slide = slide.model_copy(update={"archived": False})
            single_slide = presentation.model_copy(update={"slides": [renderable_slide]})
            return export_presentation(single_slide, profile=profile, template_path=template_path)

        _render_cached_png(target, make_pptx)
    return build_preview(presentation, context)


def render_layout_backgrounds(
    context: PreviewContext,
) -> None:
    profile = context.profile
    template_path = context.template_path
    cache_dir = context.cache_dir
    background_dir = cache_dir / profile.id / "backgrounds" / f"v{profile.version}"
    rendered_layouts: dict[int, Path] = {}
    for mapping in profile.layouts:
        target = background_dir / f"{mapping.semantic_layout}.png"
        if target.is_file():
            continue
        existing_layout = rendered_layouts.get(mapping.template_layout_index)
        if existing_layout is not None:
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(existing_layout, target)
            continue
        presentation = (
            PPTXPresentation(str(template_path))
            if template_path is not None
            else PPTXPresentation()
        )
        _delete_all_slides(presentation)
        if mapping.template_layout_index >= len(presentation.slide_layouts):
            continue
        presentation.slides.add_slide(presentation.slide_layouts[mapping.template_layout_index])

        def make_background_pptx(presentation=presentation) -> bytes:
            buffer = BytesIO()
            presentation.save(buffer)
            return buffer.getvalue()

        _render_cached_png(target, make_background_pptx)
        rendered_layouts[mapping.template_layout_index] = target


def _delete_all_slides(presentation) -> None:
    relationship_attribute = (
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    )
    slide_ids = presentation.slides._sldIdLst
    while len(slide_ids) > 0:
        relationship_id = slide_ids[0].get(relationship_attribute)
        if relationship_id is not None:
            presentation.part.drop_rel(relationship_id)
        slide_ids.remove(slide_ids[0])


def _render_pptx_png(pptx_bytes: bytes, target: Path) -> None:
    executable = _libreoffice_executable()
    if executable is None:
        raise RuntimeError("LibreOffice was not found")
    with tempfile.TemporaryDirectory() as directory:
        temporary_dir = Path(directory)
        source = temporary_dir / "slide.pptx"
        source.write_bytes(pptx_bytes)
        result = subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "png",
                "--outdir",
                str(temporary_dir),
                str(source),
            ],
            capture_output=True,
            text=True,
            timeout=30,
            check=False,
        )
        if result.returncode != 0:
            detail = result.stderr.strip() or result.stdout.strip() or "unknown error"
            raise RuntimeError(f"LibreOffice preview render failed: {detail}")
        rendered = next(temporary_dir.glob("*.png"), None)
        if rendered is None:
            raise RuntimeError("LibreOffice did not produce a preview image")
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary_target = target.with_suffix(f".{uuid4().hex}.png.tmp")
        shutil.copyfile(rendered, temporary_target)
        try:
            temporary_target.replace(target)
        finally:
            temporary_target.unlink(missing_ok=True)


def _builtin_inspection() -> dict:
    from pptx import Presentation as PPTXPresentation

    return _inspect_presentation(PPTXPresentation())


def _inspect_presentation(presentation) -> dict:
    import tempfile

    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "builtin.pptx"
        presentation.save(str(path))
        return inspect_template(path)
