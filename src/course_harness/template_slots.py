from __future__ import annotations

from pptx.enum.shapes import PP_PLACEHOLDER

TITLE_LIKE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
CONTENT_LIKE_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}


def find_placeholder_of_type(slide, placeholder_types):
    return next(
        (
            shape
            for shape in slide.placeholders
            if shape.placeholder_format.type in placeholder_types
        ),
        None,
    )


def find_title_placeholder(slide):
    return find_placeholder_of_type(slide, TITLE_LIKE_TYPES)


def find_body_placeholder(slide):
    return find_placeholder_of_type(slide, CONTENT_LIKE_TYPES)


def find_subtitle_placeholder(slide):
    return find_placeholder_of_type(slide, {PP_PLACEHOLDER.SUBTITLE})


def find_subtitle_or_body_placeholder(slide):
    return find_subtitle_placeholder(slide) or find_body_placeholder(slide)


def find_content_placeholders(slide):
    placeholders = [
        shape for shape in slide.placeholders if shape.placeholder_format.type in CONTENT_LIKE_TYPES
    ]
    placeholders.sort(key=lambda shape: shape.placeholder_format.idx)
    return placeholders


def find_slot_placeholder(slide, slot_mappings: dict[str, int], slot: str):
    placeholder_index = slot_mappings.get(slot)
    if placeholder_index is None:
        return None
    return next(
        (
            shape
            for shape in slide.placeholders
            if shape.placeholder_format.idx == placeholder_index
        ),
        None,
    )


def find_slot_placeholder_or(slide, slot_mappings: dict[str, int], slot: str, fallback):
    return find_slot_placeholder(slide, slot_mappings, slot) or fallback(slide)


def find_column_placeholder_groups(slide) -> list[list]:
    placeholders = [
        shape for shape in slide.placeholders if shape.placeholder_format.type in CONTENT_LIKE_TYPES
    ]
    placeholders.sort(key=lambda shape: (shape.left, shape.top))
    groups: list[list] = []
    for placeholder in placeholders:
        if not groups or abs(placeholder.left - groups[-1][0].left) > 500000:
            groups.append([placeholder])
        else:
            groups[-1].append(placeholder)
    return groups
